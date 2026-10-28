"""WP-2 import / export: the ASCII Import Wizard options and import filters, JCAMP-DX (all
ASDF compression forms from the IUPAC specification's character tables, Y-check, NTUPLES,
PEAK TABLE, round trip), Thermo/Galactic SPC (fixtures built from the published header
layout), MATLAB, NetCDF 3, HDF5 (optional h5py), WAV, image → matrix, Excel (all sheets),
and CSV / XLSX / report export. All fixtures are made by the tests (no third-party data)."""
from __future__ import annotations

import struct

import numpy as np
import pytest

from rsstatchem.core.analysis import exporters as ex
from rsstatchem.core.analysis import importers as imp
from rsstatchem.core.analysis.importers import ascii_wizard as aw
from rsstatchem.core.analysis.importers import jcamp, spc
from rsstatchem.ui.origin import worksheet_ops as wo
from rsstatchem.ui.origin.model import ReportSheet, Workbook


# ====================================================================== ASCII wizard
def _write(path, text, enc="utf-8"):
    path.write_text(text, encoding=enc)
    return path


def test_ascii_auto_subheaders_long_name_units_comments(tmp_path):
    p = _write(tmp_path / "s.csv", "Time,Signal,Ref\ns,mV,mV\nrun 1,run 1,blank\n"
                                   "0,1.5,0.1\n1,2.5,0.2\n2,3.5,0.3\n")
    t = aw.read(p)
    assert [c.long_name for c in t.columns] == ["Time", "Signal", "Ref"]
    assert [c.units for c in t.columns] == ["s", "mV", "mV"]
    assert [c.comments for c in t.columns] == ["run 1", "run 1", "blank"]
    assert [c.designation for c in t.columns] == ["X", "Y", "Y"]
    assert t.columns[1].values == [1.5, 2.5, 3.5]
    assert t.meta["delimiter"] == "," and t.meta["subheader_lines"] == 3


def test_ascii_semicolon_decimal_comma_main_header_and_params(tmp_path):
    text = ("Instrument: UV-1800\nOperator = RS\n\nWavelength;Absorbance\n"
            "200,5;0,123\n201,0;1,5E-01\n201,5;-0,02\n")
    p = _write(tmp_path / "uv.txt", text, "cp1252")
    o = aw.AsciiOptions(skip_lines=3)
    t = aw.read(p, o)
    assert t.meta["delimiter"] == ";" and t.meta["decimal"] == ","
    assert t.columns[0].values == [200.5, 201.0, 201.5]
    assert t.columns[1].values == pytest.approx([0.123, 0.15, -0.02])
    assert t.meta["params"] == {"Instrument": "UV-1800", "Operator": "RS"}


def test_ascii_whitespace_partial_import_and_designation_pattern(tmp_path):
    rows = "\n".join(f"{i} {i * 2} {i * 3} {i * 4}" for i in range(1, 21))
    p = _write(tmp_path / "w.dat", "# comment\n" + rows + "\n")
    o = aw.AsciiOptions(comment_char="#", header_lines=0, first_row=3, last_row=12,
                        read_every=3, columns="1,3:4", designations="XYXY")
    t = aw.read(p, o)
    assert t.meta["delimiter"] == "whitespace"
    assert [c.designation for c in t.columns] == ["X", "Y", "X"]
    assert t.columns[0].values == [3.0, 6.0, 9.0, 12.0]
    assert t.columns[2].values == [12.0, 24.0, 36.0, 48.0]
    assert aw.parse_column_list("2:*", 5) == [1, 2, 3, 4]


def test_ascii_dates_text_and_fortran_exponents(tmp_path):
    p = _write(tmp_path / "d.csv", "Date,Value,Label\n2024-01-01,1.0D+02,a\n"
                                   "2024-01-02,2.5d-01,b\n")
    t = aw.read(p, aw.AsciiOptions(dates="julian"))
    assert t.columns[0].values == [2460310.5, 2460311.5]          # Julian day numbers
    assert t.columns[1].values == [100.0, 0.25]
    assert t.columns[2].values == ["a", "b"]
    t2 = aw.read(p, aw.AsciiOptions(dates="excel"))
    assert t2.columns[0].values == [45292.0, 45293.0]             # Excel serial days


def test_import_filter_save_list_apply_by_pattern(tmp_path):
    folder = tmp_path / "filters"
    o = aw.AsciiOptions(delimiter=";", decimal=",", header_lines=1, designations="")
    aw.save_filter("Lab UV", o, "*.uvx", folder)
    assert [f["name"] for f in aw.list_filters(folder)] == ["Lab UV"]
    assert aw.load_filter("Lab UV", folder).delimiter == ";"
    name, opts = aw.find_filter(tmp_path / "sample.UVX", folder)
    assert name == "Lab UV" and opts.decimal == ","
    assert aw.find_filter(tmp_path / "x.csv", folder) is None
    assert aw.delete_filter("Lab UV", folder) and aw.list_filters(folder) == []


# ====================================================================== JCAMP-DX
# The same ordinates 1 2 3 3 2 1 0 -1 -2 -3 (abscissa 1) in every ASDF form, written from
# the specification's tables: SQZ @=0, A-I=1..9, a-i=-1..-9; DIF %=0, J-R=1..9, j-r=-1..-9;
# DUP S-Z=1..8, s=9 (repeat count of the preceding value or difference).
FORMS = {"AFFN": "1 1 2 3 3 2 1 0 -1 -2 -3",
         "PAC": "1+1+2+3+3+2+1+0-1-2-3",
         "SQZ": "1ABCCBA@abc",
         "DIF": "1AJJ%jjjjjj",
         "DIFDUP": "1AJT%jX"}
EXPECTED = [1, 2, 3, 3, 2, 1, 0, -1, -2, -3]


@pytest.mark.parametrize("form", list(FORMS))
def test_jcamp_asdf_forms_decode_identically(form):
    vals, is_dif = jcamp.decode_line(FORMS[form])
    assert vals[0] == 1 and vals[1:] == EXPECTED
    assert is_dif == form.startswith("DIF")


def test_jcamp_multi_digit_and_dup_of_values():
    # SQZ with several digits, DUP of an absolute value, DIF with several digits
    assert jcamp.decode_line("100A23B5U")[0] == [100, 123, 25, 25, 25]
    assert len(jcamp.decode_line("1AS2")[0]) == 1 + 12   # DUP "S2" = 12 copies
    assert jcamp.decode_line("0@J10K25")[0] == [0, 0, 110, 335]
    assert jcamp.decode_line("5 1.5E+02,-2.5e-1")[0] == [5, 150.0, -0.25]


def _jdx(body: str, **ldr) -> str:
    head = {"TITLE": "test", "JCAMP-DX": "4.24", "DATA TYPE": "INFRARED SPECTRUM",
            "XUNITS": "1/CM", "YUNITS": "ABSORBANCE", "XFACTOR": "1", "YFACTOR": "1"}
    head.update(ldr)
    return "\n".join(f"##{k}={v}" for k, v in head.items()) + "\n" + body + "\n##END=\n"


def test_jcamp_xydata_dif_y_check_factors_and_x_axis():
    # two DIF lines: the last ordinate of line 1 (Y-check) repeats as the first of line 2
    body = "##XYDATA=(X++(Y..Y))\n1000A10NN%\n1003A20j0j0j0\n"
    text = _jdx(body, FIRSTX="1000", LASTX="1006", NPOINTS="7", YFACTOR="0.5",
                **{"$$ note": "x"})
    b = jcamp.read_text(text)[0]
    assert b["kind"] == "xydata"
    np.testing.assert_allclose(b["x"], [1000, 1001, 1002, 1003, 1004, 1005, 1006])
    np.testing.assert_allclose(b["ys"][0][1], np.array([110, 115, 120, 120, 110, 100, 90])
                               * 0.5)
    with pytest.raises(imp.ImportFailed, match="Y-check"):
        jcamp.read_text(_jdx("##XYDATA=(X++(Y..Y))\n1000A10NN%\n1003A21j0j0j0\n",
                             FIRSTX="1000", LASTX="1006", NPOINTS="7"))


def test_jcamp_xypoints_peak_table_and_link_blocks():
    text = ("##TITLE=link\n##JCAMP-DX=5.01\n##BLOCKS=2\n" +
            _jdx("##XYPOINTS=(XY..XY)\n1.0, 10.0; 2.0, 20.0\n3.0, 30.0\n", TITLE="pts",
                 XFACTOR="2") +
            _jdx("##PEAK TABLE=(XYW..XYW)\n100,5,1 200,7,2\n", TITLE="peaks",
                 **{"DATA TYPE": "MASS SPECTRUM", "XUNITS": "M/Z"}) + "##END=\n")
    blocks = jcamp.read_text(text)
    assert [b["kind"] for b in blocks] == ["xypoints", "peaks"]
    np.testing.assert_allclose(blocks[0]["x"], [2, 4, 6])
    np.testing.assert_allclose(blocks[1]["ys"][0][1], [5, 7])
    tables = jcamp.tables(blocks, "f")
    assert tables[1].columns[0].long_name == "m/z" and tables[1].columns[2].long_name == "W"


def test_jcamp_ntuples_nmr_real_imaginary_pages():
    text = """##TITLE=nmr fid
##JCAMP-DX=5.00
##DATA TYPE=NMR FID
##NTUPLES=NMR FID
##VAR_NAME=TIME, FID/REAL, FID/IMAG, PAGE NUMBER
##SYMBOL=X, R, I, N
##VAR_TYPE=INDEPENDENT, DEPENDENT, DEPENDENT, PAGE
##VAR_FORM=AFFN, ASDF, ASDF, AFFN
##VAR_DIM=4, 4, 4, 2
##UNITS=SECONDS, ARBITRARY UNITS, ARBITRARY UNITS,
##FIRST=0.0, 1, 0, 1
##LAST=0.3, 4, 0, 2
##FACTOR=0.1, 2.0, 1.0, 1
##PAGE=N=1
##DATA TABLE=(X++(R..R)), XYDATA
0AJ%J
##PAGE=N=2
##DATA TABLE=(X++(I..I)), XYDATA
0@AB@
##END NTUPLES=NMR FID
##END=
"""
    b = jcamp.read_text(text)[0]
    assert b["kind"] == "ntuples" and [p["name"] for p in b["pages"]] == ["FID/REAL",
                                                                         "FID/IMAG"]
    np.testing.assert_allclose(b["pages"][0]["x"], [0.0, 0.1, 0.2, 0.3])
    np.testing.assert_allclose(b["pages"][0]["y"], [2, 4, 4, 6])
    np.testing.assert_allclose(b["pages"][1]["y"], [0, 1, 2, 0])
    t = jcamp.tables([b], "fid")[0]
    assert [c.long_name for c in t.columns] == ["Time", "FID/REAL", "FID/IMAG"]


@pytest.mark.parametrize("compression", ["difdup", "affn"])
def test_jcamp_writer_round_trip(tmp_path, compression):
    x = np.linspace(4000, 400, 901)
    y = np.exp(-((x - 1650) / 25) ** 2) + 0.3 * np.exp(-((x - 2900) / 60) ** 2)
    p = jcamp.write(tmp_path / f"r_{compression}.jdx", x, y, "round trip",
                    compression=compression, per_line=8)
    t = imp.read_file(p)[0]
    assert t.columns[0].long_name == "Wavenumber" and t.columns[0].units == "cm⁻¹"
    np.testing.assert_allclose(np.array(t.columns[0].values), x, atol=1e-9)
    np.testing.assert_allclose(np.array(t.columns[1].values), y, atol=1e-7)


def test_jcamp_cross_check_with_optional_jcamp_package(tmp_path):
    jc = pytest.importorskip("jcamp")
    x = np.linspace(500, 1500, 101)
    y = np.sin(x / 50) + 2
    p = jcamp.write(tmp_path / "c.jdx", x, y)
    ref = jc.jcamp_readfile(str(p))
    np.testing.assert_allclose(ref["y"], y, atol=1e-6)


# ====================================================================== SPC
def _spc_header(flags, npts, first, last, nsub, exp=-128, xtype=1, ytype=2, version=0x4B,
                logoff=0, catxt=b"", cmnt=b"made by test", zinc=0.0):
    return struct.pack(spc._MAIN_NEW, flags, version, 0, exp, npts, first, last, nsub, xtype,
                       ytype, 0, 0, 0, b"", b"", 0, b"", cmnt.ljust(130, b"\0"),
                       catxt.ljust(30, b"\0"), logoff, 0, 0, 0, 0, 0.0, b"", zinc, 0, 0.0,
                       0, b"")


def _sub(npts=0, exp=-128, z=0.0, index=0):
    return struct.pack(spc._SUB, 0, exp, index, z, 0.0, 0.0, npts, 0, 0.0, b"")


def test_spc_new_format_single_float(tmp_path):
    y = np.array([0.1, 0.5, 1.0, 0.5, 0.1], np.float32)
    data = _spc_header(0, 5, 400.0, 404.0, 1) + _sub() + y.tobytes()
    s = spc.read_bytes(data)
    np.testing.assert_allclose(s["x"][0], [400, 401, 402, 403, 404])
    np.testing.assert_allclose(s["y"][0], y)
    assert s["xlabel"] == "Wavenumber" and s["ylabel"] == "Absorbance"
    assert s["comment"] == "made by test"


def test_spc_scaled_integers_multifile_and_z(tmp_path):
    # 32-bit integers scaled by 2^(exp-32); each subfile has its own exponent (TMULTI)
    y1 = np.array([1, 2, 3, 4], np.int32) * 2 ** 20
    y2 = np.array([4, 3, 2, 1], np.int32) * 2 ** 24
    flags = spc.TMULTI
    data = _spc_header(flags, 4, 0.0, 3.0, 2, exp=0, zinc=0.5) + \
        _sub(exp=12, z=0.0) + y1.tobytes() + _sub(exp=8, z=0.5, index=1) + y2.tobytes()
    s = spc.read_bytes(data)
    np.testing.assert_allclose(s["y"][0], np.array([1, 2, 3, 4]) * 2 ** 20 * 2.0 ** (12 - 32))
    np.testing.assert_allclose(s["y"][1], np.array([4, 3, 2, 1]) * 2 ** 24 * 2.0 ** (8 - 32))
    assert s["z"] == [0.0, 0.5]
    t = spc.tables(s, "multi")[0]
    assert [c.designation for c in t.columns] == ["X", "Y", "Y"]
    assert t.columns[2].comments.endswith("= 0.5")


def test_spc_16bit_txvals_and_txyxys_and_labels():
    x = np.array([1.0, 2.0, 4.0], np.float32)
    y = np.array([100, -200, 300], np.int16)
    data = _spc_header(spc.TSPREC | spc.TXVALS | spc.TALABS, 3, 1.0, 4.0, 1, exp=16,
                       catxt=b"Time (s)\0Counts\0") + x.tobytes() + _sub(exp=16) + y.tobytes()
    s = spc.read_bytes(data)
    np.testing.assert_allclose(s["x"][0], [1, 2, 4])
    np.testing.assert_allclose(s["y"][0], [100, -200, 300])
    assert s["xlabel"] == "Time (s)" and s["ylabel"] == "Counts"
    # XY-XY multifile: every subfile has its own X array and point count
    flags = spc.TMULTI | spc.TXVALS | spc.TXYXYS
    xa, ya = np.array([1, 2], np.float32), np.array([5, 6], np.float32)
    xb, yb = np.array([1, 2, 3], np.float32), np.array([7, 8, 9], np.float32)
    data = _spc_header(flags, 0, 0.0, 0.0, 2) + _sub(2) + xa.tobytes() + ya.tobytes() + \
        _sub(3) + xb.tobytes() + yb.tobytes()
    s = spc.read_bytes(data)
    assert [len(v) for v in s["x"]] == [2, 3]
    t = spc.tables(s)[0]
    assert [c.designation for c in t.columns] == ["X", "Y", "X", "Y"]


def test_spc_log_block_and_old_format():
    y = np.array([0.25, 0.5], np.float32)
    head = _spc_header(0, 2, 0.0, 1.0, 1, logoff=512 + 32 + 8)
    text = b"Operator=RS\r\nSample=water\r\n\0"
    log = struct.pack("<IIII", 64 + len(text), 0, 64, 0) + b"\0" * 48 + text
    s = spc.read_bytes(head + _sub() + y.tobytes() + log)
    assert s["log"] == {"Operator": "RS", "Sample": "water"}
    # old format 0x4D: 32-bit integers with swapped 16-bit words, first subheader at 224
    vals = np.array([65537 * 3, 7], np.int64)      # 0x00030003 and 0x00000007
    words = b"".join(struct.pack("<HH", (int(v) >> 16) & 0xFFFF, int(v) & 0xFFFF)
                     for v in vals)
    old = bytearray(256)
    struct.pack_into("<BBhfffBB", old, 0, 0, 0x4D, 32, 2.0, 10.0, 20.0, 3, 0)
    s = spc.read_bytes(bytes(old) + words)
    np.testing.assert_allclose(s["x"][0], [10, 20])
    np.testing.assert_allclose(s["y"][0], vals)
    with pytest.raises(imp.ImportFailed):
        spc.read_bytes(b"\0\x99" + b"\0" * 600)


# ====================================================================== other readers
def test_matlab_netcdf_wav_image_excel(tmp_path):
    from scipy import io as sio
    sio.savemat(tmp_path / "m.mat", {"t": np.arange(5.0), "y": np.arange(5.0) ** 2,
                                     "M": np.arange(12.0).reshape(3, 4) * 1.0,
                                     "big": np.ones((80, 90)),
                                     "s": {"a": np.array([1.0, 2.0])}})
    tables = imp.read_file(tmp_path / "m.mat")
    sheet = next(t for t in tables if t.kind == "sheet" and t.column("y") is not None)
    assert sheet.column("y").values == [0, 1, 4, 9, 16]
    mats = [t for t in tables if t.kind == "matrix"]
    assert any(t.matrix.shape == (80, 90) for t in mats)
    assert any(t.column("s.a") is not None for t in tables)

    from scipy.io import netcdf_file
    f = netcdf_file(str(tmp_path / "n.nc"), "w")
    f.createDimension("time", 4)
    tv = f.createVariable("time", "d", ("time",))
    tv[:] = [0, 1, 2, 3]
    tv.units = b"s"
    v = f.createVariable("temp", "d", ("time",))
    v[:] = [20.0, 21.0, 22.5, 23.0]
    v.units = b"degC"
    v.long_name = b"Temperature"
    f.close()
    t = imp.read_file(tmp_path / "n.nc")[0]
    assert t.columns[0].long_name == "time" and t.columns[0].designation == "X"
    assert t.columns[1].long_name == "Temperature" and t.columns[1].units == "degC"

    from scipy.io import wavfile
    wavfile.write(tmp_path / "a.wav", 8000, (np.sin(np.arange(80) / 5) * 3000)
                  .astype(np.int16))
    t = imp.read_file(tmp_path / "a.wav")[0]
    assert t.columns[0].units == "s" and t.columns[0].values[1] == pytest.approx(1 / 8000)
    assert max(abs(v) for v in t.columns[1].values) < 1.0

    from PIL import Image
    Image.fromarray(np.arange(12, dtype=np.uint8).reshape(3, 4) * 20).save(tmp_path / "i.png")
    t = imp.read_file(tmp_path / "i.png")[0]
    assert t.kind == "matrix" and t.matrix.shape == (3, 4) and t.matrix[2, 3] == 220

    import pandas as pd
    with pd.ExcelWriter(tmp_path / "x.xlsx") as w:
        pd.DataFrame({"a": [1, 2], "b": [3, 4]}).to_excel(w, sheet_name="One", index=False)
        pd.DataFrame({"c": ["u", "v"]}).to_excel(w, sheet_name="Two", index=False)
    ts = imp.read_file(tmp_path / "x.xlsx")
    assert [t.name for t in ts] == ["One", "Two"] and ts[1].columns[0].values == ["u", "v"]


def test_hdf5_needs_optional_h5py(tmp_path):
    from rsstatchem.core.analysis.importers import hdf
    p = tmp_path / "d.h5"
    if not hdf.have_h5py():
        p.write_bytes(hdf.HDF5_MAGIC + b"\0" * 100)
        with pytest.raises(imp.ImportFailed, match="h5py"):
            imp.read_file(p)
        return
    import h5py
    with h5py.File(p, "w") as f:
        d = f.create_dataset("group/x", data=np.arange(6.0))
        d.attrs["units"] = "nm"
        f.create_dataset("group/y", data=np.arange(6.0) * 2)
        f.create_dataset("img", data=np.ones((70, 80)))
    tables = imp.read_file(p)
    sheet = tables[0]
    assert sheet.column("group.x").units == "nm"
    assert any(t.kind == "matrix" and t.matrix.shape == (70, 80) for t in tables)


# ====================================================================== tables → books
def test_tables_to_books_modes():
    t1 = imp.ImportedTable("a", [imp.ImportColumn("X", [1.0, 2.0], designation="X"),
                                 imp.ImportColumn("Y", [3.0, 4.0])])
    t2 = imp.ImportedTable("b", [imp.ImportColumn("X", [1.0, 2.0, 3.0], designation="X"),
                                 imp.ImportColumn("Y", [5.0, 6.0, 7.0])])
    names = iter(f"Book{k}" for k in range(1, 99))
    books = wo.tables_to_books([("f1", [t1]), ("f2", [t2])], "books", lambda b: next(names))
    assert len(books) == 2 and books[1].long_name == "f2"
    names = iter(f"Book{k}" for k in range(1, 99))
    (book,) = wo.tables_to_books([("f1", [t1]), ("f2", [t2])], "sheets", lambda b: next(names))
    assert [s.name for s in book.sheets] == ["f1", "f2"]
    names = iter(f"Book{k}" for k in range(1, 99))
    (book,) = wo.tables_to_books([("f1", [t1]), ("f2", [t2])], "columns",
                                 lambda b: next(names))
    ws = book.sheets[0]
    assert len(ws.columns) == 4 and ws.columns[3].numeric(3).tolist() == [5, 6, 7]
    assert ws.columns[2].comments.endswith("f2")


# ====================================================================== export
def _sheet():
    book = Workbook.new("Book1", 3, 4)
    ws = book.sheets[0]
    ws.set_column_values(0, [1, 2, 3, 4])
    ws.set_column_values(1, [1.5, 2.5, None, "n/a"])
    ws.set_column_values(2, [10, 20, 30, 40])
    ws.columns[1].long_name, ws.columns[1].units = "Signal", "mV"
    ws.mask_cells(2, [1])
    return book, ws


def test_export_csv_headers_decimal_and_masked(tmp_path):
    _, ws = _sheet()
    p = ex.write_csv(ws, tmp_path / "o.csv", ";", ",", ("name", "long_name", "units"),
                     masked="blank")
    lines = p.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "A;B;C" and lines[1] == "A;Signal;C" and lines[2] == ";mV;"
    assert lines[3] == "1;1,5;10" and lines[4] == "2;2,5;"        # masked C2 left empty
    assert lines[6] == "4;n/a;40"
    back = aw.read(p, aw.AsciiOptions(skip_lines=1, header_lines=2, units_line=2,
                                      comments_from=0))
    assert back.columns[1].units == "mV" and back.columns[1].values[:2] == [1.5, 2.5]


def test_export_xlsx_multiple_sheets_round_trip(tmp_path):
    book, ws = _sheet()
    ws2 = ws.copy()
    ws2.name = "Copy"
    p = ex.write_xlsx([("Sheet1", ws), ("Sheet1", ws2)], tmp_path / "o.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert wb.sheetnames == ["Sheet1", "Sheet1_2"]
    rows = list(wb["Sheet1"].values)
    assert rows[0] == ("A", "Signal", "C") and rows[1] == (None, "mV", None)
    assert rows[2] == (1, 1.5, 10) and wb["Sheet1"]["A1"].font.bold


def test_export_report_docx_and_html(tmp_path):
    rep = ReportSheet("Report1", "My report")
    rep.add_table("Parameters", ["Name", "Value"], [["a", "1.0"], ["b", "2.0"]], ["note"])
    rep.add_text("Notes", "Some text")
    p = ex.write_report_docx(rep, tmp_path / "r.docx")
    from docx import Document
    doc = Document(p)
    assert doc.tables[0].cell(1, 0).text == "a"
    assert any("Some text" in par.text for par in doc.paragraphs)
    h = ex.write_html("<html>x</html>", tmp_path / "r.html")
    assert h.read_text() == "<html>x</html>"
