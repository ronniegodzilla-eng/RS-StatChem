"""WP6: Output documents (.rso round trip), Excel export, APA TableLook, pivot editing."""
from __future__ import annotations

import zipfile

import pytest

from rsstatchem.core import pivot as pv
from rsstatchem.core import report
from rsstatchem.core.output import Image, OutputDocument, Section, Table, Text
from rsstatchem.core.output_store import load_output, save_output
from rsstatchem.core.report_xlsx import export_xlsx, parse_number


def _png() -> bytes:
    import io

    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    fig = Figure(figsize=(2, 1.5), dpi=60)
    FigureCanvasAgg(fig)
    fig.add_subplot().plot([1, 2, 3], [2, 1, 3])
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    return buf.getvalue()


def _group_stats() -> Table:
    return Table("Group Statistics", ["", "Group", "N", "Mean", "Std. Deviation"],
                 [["Score", "Control", "10", "12.500", ".845"],
                  ["", "Treatment", "12", "14.250", "1.020"],
                  ["Age", "Control", "10", "20.100", "2.000"],
                  ["", "Treatment", "12", "21.300", "1.500"]],
                 row_header_cols=2, footnotes=["a. Example footnote."])


def _wide() -> Table:
    return Table("Independent Samples Test",
                 ["", ("Levene", "F"), ("Levene", "Sig."), ("t-test", "t"), ("t-test", "df"),
                  ("t-test", "Sig. (2-tailed)")],
                 [["Equal variances assumed", "1.234", ".280", "-2.110", "20", ".048"],
                  ["Equal variances not assumed", "", "", "-2.090", "18.611", "<.001"]])


@pytest.fixture()
def doc():
    d = OutputDocument()
    d.add(Section("T-Test", [Text("", "T-TEST GROUPS=g(1 2).", style="log"), _group_stats(),
                             _wide(), Image("Chart", _png(), 300),
                             Text("Interpretation", "The difference is significant.")]))
    d.add(Section("Notes", [Text("", "Inserted heading", style="title")]))
    return d


def test_rso_round_trip_keeps_everything(tmp_path, doc):
    t = doc.sections[0].items[1]
    t.hidden_rows = [3]
    path = tmp_path / "out.rso"
    save_output(doc, path, "apa")
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
    assert "output.json" in names and any(n.startswith("images/") for n in names)
    back, look = load_output(path)
    assert look == "apa"
    assert [s.title for s in back.sections] == ["T-Test", "Notes"]
    items = back.sections[0].items
    assert [type(i).__name__ for i in items] == ["Text", "Table", "Table", "Image", "Text"]
    assert items[1].rows == t.rows and items[1].columns == t.columns
    assert items[1].hidden_rows == [3] and items[1].footnotes == t.footnotes
    assert items[2].columns == _wide().columns  # (group, sub) tuples survive JSON
    assert items[3].png == doc.sections[0].items[3].png and items[3].width == 300
    assert items[0].style == "log" and back.sections[1].items[0].style == "title"
    assert len({i.id for s in back.sections for i in s.items}) == 6


def test_rso_rejects_other_files(tmp_path):
    bad = tmp_path / "x.rso"
    with zipfile.ZipFile(bad, "w") as zf:
        zf.writestr("output.json", '{"format": "something"}')
    with pytest.raises(ValueError):
        load_output(bad)


def test_excel_export_one_sheet_per_table(tmp_path, doc):
    from openpyxl import load_workbook

    path = tmp_path / "out.xlsx"
    assert export_xlsx(doc, str(path), "sheets") == 2
    wb = load_workbook(path)
    assert wb.sheetnames[:2] == ["1 Group Statistics", "2 Independent Samples Test"]
    assert any(n.startswith("Chart") for n in wb.sheetnames) and "Notes" in wb.sheetnames
    ws = wb["1 Group Statistics"]
    values = [tuple(row) for row in ws.iter_rows(values_only=True)]
    assert values[1][0] == "Group Statistics"
    assert values[2][1:] == ("Group", "N", "Mean", "Std. Deviation")
    assert values[3] == ("Score", "Control", 10, 12.5, 0.845)
    assert ws.cell(4, 5).number_format == "#.000"     # ".845" keeps its look
    assert ws.cell(4, 4).number_format == "0.000"
    assert ws.cell(3, 1).font.bold and ws.cell(3, 2).fill.fgColor.rgb.endswith("F2F2F2")
    ws2 = wb["2 Independent Samples Test"]
    assert ws2.cell(3, 2).value == "Levene" and "B3:C3" in {str(r) for r in
                                                              ws2.merged_cells.ranges}
    assert ws2.cell(6, 6).value == "<.001"          # a bound stays text
    assert len(wb[[n for n in wb.sheetnames if n.startswith("Chart")][0]]._images) == 1


def test_excel_export_single_sheet_apa(tmp_path, doc):
    from openpyxl import load_workbook

    path = tmp_path / "apa.xlsx"
    export_xlsx(doc, str(path), "single", look="apa")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Output"]
    col_a = [r[0] for r in wb["Output"].iter_rows(values_only=True)]
    assert "Table 1" in col_a and "Table 2" in col_a and "Group Statistics" in col_a
    assert "Note. a. Example footnote." in col_a
    assert "The difference is significant." in col_a


def test_parse_number_formats():
    assert parse_number(".045") == (0.045, "#.000")
    assert parse_number("-1.50") == (-1.5, "0.00")
    assert parse_number("12") == (12.0, "0")
    assert parse_number("45.5%") == (0.455, "0.0%")
    assert parse_number("<.001") is None and parse_number("1.2ᵃ") is None


def test_apa_look_in_html_word_and_docx(tmp_path, doc):
    html = report.document_html(doc, look="apa")
    assert "Table 1" in html and "Table 2" in html and "<i>Note.</i>" in html
    assert "Times New Roman" in html
    spss = report.document_html(doc)
    assert "Table 1" not in spss and "Times New Roman" not in spss
    word = report.clipboard_html([doc.sections[0].items[2]], look="apa", number=2)
    assert "Table 2" in word and "font-style:italic" in word
    path = tmp_path / "apa.docx"
    report.export_docx(doc, str(path), look="apa")
    from docx import Document

    text = [p.text for p in Document(str(path)).paragraphs]
    assert "Table 1" in text and "Table 2" in text and any(t.startswith("Note.") for t in text)


def test_pivot_moves_dimensions_and_round_trips():
    t = _group_stats()
    cube = pv.cube(t)
    assert cube.values[0] == ["Score", "Age"] and cube.values[2][0] == "N"
    wide = pv.pivot(t, rows=[1], cols=[0, 2])
    assert wide.columns[1] == ("Score", "N") and wide.rows[0][:2] == ["Control", "10"]
    back = pv.pivot(wide, rows=[0, 1], cols=[2])
    assert back.rows == t.rows and back.columns == t.columns
    layer = pv.pivot(t, rows=[1], cols=[2], layers={0: "Age"})
    assert layer.rows == [["Control", "10", "20.100", "2.000"],
                          ["Treatment", "12", "21.300", "1.500"]]
    assert "Age" in layer.title
    assert pv.reset_pivot(layer).rows == t.rows
    with pytest.raises(ValueError):
        pv.pivot(t, rows=[0], cols=[2])  # a dimension left out


def test_hide_rows_columns_and_decimals():
    t = _group_stats()
    t.hidden_rows = [0]
    t.hidden_cols = [4]
    shown = pv.visible(t)
    assert shown.rows[0] == ["Score", "Treatment", "12", "14.250"]  # label moved down
    assert len(shown.columns) == 4
    assert "Std. Deviation" not in report.table_html(t)
    assert pv.set_decimals(t, 1, cols=[3, 4]) == 8
    assert t.rows[0][3:] == ["12.5", ".8"] and t.rows[0][2] == "10"
    assert pv.format_cell("<.001", 1) == "<.001" and pv.format_cell("45.25%", 1) == "45.2%"


def test_viewer_save_open_edit_and_pivot(qapp, tmp_path, doc):
    from PySide6.QtCore import QSettings

    from rsstatchem.ui.output_viewer import OutputViewer

    v = OutputViewer(QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat))
    for s in doc.sections:
        v.add_section(s)
    assert v.dirty
    table = v.doc.sections[0].items[1]
    v.select(f"i{table.id}")
    new = pv.pivot(table, rows=[1], cols=[0, 2])
    v.replace_item(table, new)
    assert v.doc.sections[0].items[1] is new and v.selected == f"i{new.id}"
    v.add_text("Hasil uji t", "title")
    path = str(tmp_path / "viewer.rso")
    assert v.save_to(path) and not v.dirty and "viewer.rso" in v.windowTitle()
    v.set_look("apa")
    assert "Table 1" in v.view.toPlainText()
    v2 = OutputViewer(QSettings(str(tmp_path / "s2.ini"), QSettings.Format.IniFormat))
    assert v2.open_path(path)
    assert [s.title for s in v2.doc.sections] == ["T-Test", "Notes"]
    assert v2.root.childCount() == 2 and "Hasil uji t" in v2.view.toPlainText()
    assert v2.doc.sections[0].items[1].pivot is not None
    v2.select(f"i{v2.doc.sections[0].items[1].id}")
    v2.reset_selected()
    assert v2.doc.sections[0].items[1].rows == _group_stats().rows


def test_table_editor_dialog_hides_and_edits(qapp):
    from PySide6.QtCore import Qt

    from rsstatchem.ui.pivot_dialogs import PivotTraysDialog, TableEditorDialog

    t = _group_stats()
    dlg = TableEditorDialog(t)
    dlg.title_edit.setText("Statistik Kelompok")
    dlg.grid.item(1, 2).setText("11")             # first data row, N column
    dlg.row_list.item(3).setCheckState(Qt.CheckState.Unchecked)
    dlg.notes.setPlainText("Catatan.")
    out = dlg.result_table()
    assert out.title == "Statistik Kelompok" and out.rows[0][2] == "11"
    assert out.hidden_rows == [3] and out.footnotes == ["Catatan."]
    assert t.rows[0][2] == "10"                   # original untouched
    trays = PivotTraysDialog(t)
    trays.trays["rows"].setCurrentRow(0)
    trays._move("rows", "layers")
    spec = trays.layout_spec()
    assert spec["layers"] == {0: "Score"} and spec["rows"] == [1]
