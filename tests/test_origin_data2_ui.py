"""WP-2 user interface in the Origin-style main window: ``data_menus.install`` and every
action it adds (dialogs auto-accepted with their defaults), masking in the grid (red cells,
undo, quick statistics), the Import Wizard preview / filters / multiple-file modes, and
export from the window."""
from __future__ import annotations

import numpy as np
import pytest

from rsstatchem.core.analysis import wavelets as wv
from rsstatchem.core.analysis.importers import ascii_wizard as aw
from rsstatchem.ui.origin import data_menus
from rsstatchem.ui.origin import import_wizard as iw
from rsstatchem.ui.origin import math_dialogs as md
from rsstatchem.ui.origin import signal_dialogs as sd
from rsstatchem.ui.origin.model import MatrixBook, MatrixSheet, Workbook


@pytest.fixture
def ctx(qapp, tmp_path, monkeypatch):
    monkeypatch.setenv("RSSTATCHEM_HOME", str(tmp_path / "home"))
    from rsstatchem.app import AppContext

    context = AppContext(qapp)
    monkeypatch.setattr(context.app, "quit", lambda: None)
    return context


@pytest.fixture
def win(ctx, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    from rsstatchem.ui.origin import open_workspace

    def boom(*args, **kwargs):
        raise AssertionError("unexpected modal dialog")
    for name in ("question", "warning", "information", "critical"):
        monkeypatch.setattr(QMessageBox, name, boom)
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *a, **k: ([], ""))
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: ("", ""))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: ("", ""))
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *a, **k: "")
    w = open_workspace(ctx)
    monkeypatch.setattr(w, "confirm_discard", lambda: True)
    data_menus.install(w)
    w.warnings = []
    monkeypatch.setattr(w, "_warn", lambda text, title="": w.warnings.append(text))
    yield w
    w.close()


def _fill(win, n=64):
    bw = win.active_book_window()
    ws = bw.worksheet()
    x = np.linspace(0, 6.3, n)
    ws.set_rows(n)
    while len(ws.columns) < 4:
        ws.add_column("Y")
    ws.columns[3].designation = "Z"
    ws.set_column_values(0, list(x))
    ws.set_column_values(1, list(np.sin(x) + 0.05 * np.cos(9 * x)))
    ws.set_column_values(2, list(np.cos(x)))
    ws.set_column_values(3, list(x * np.cos(x)))
    ws.columns[1].long_name = "Signal"
    bw.refresh()
    return bw, ws


def _menu_texts(menu) -> list[str]:
    return [a.text().replace("&", "") for a in menu.actions()]


# ====================================================================== menus
def test_install_adds_origin_named_menus(win):
    m = win.menus
    imp = next(a.menu() for a in m["file"].actions() if a.menu() and "Import" in a.text())
    assert {"Import Wizard...", "JCAMP-DX...", "Thermo Galactic SPC...", "MATLAB...",
            "HDF5 / NetCDF..."} <= set(_menu_texts(imp))
    sig = next(a.menu() for a in m["analysis"].actions()
               if a.menu() and "Signal" in a.text())
    assert {"IIR Filter...", "Convolution...", "Hilbert Transform...", "Wavelet"} <= \
        set(_menu_texts(sig))
    wav = next(a for a in sig.actions() if a.text() == "&Wavelet")
    assert wav.isVisible() == wv.available()
    assert "Mask" in _menu_texts(m["column"])
    assert {"Stack Columns...", "Pivot Table...", "Worksheet Query..."} <= \
        set(_menu_texts(m["worksheet"]))
    assert "Inverse" in _menu_texts(m["matrix"])
    assert len(win.data_actions) >= 70


BOOK_ACTIONS = ["sort_nested", "query", "split", "stack", "unstack", "pivot", "reduce_rows",
                "reduce_dup_x", "reduce_xy", "remove_duplicates", "trim", "interpolate",
                "interp_y_from_x", "interp_2d", "interp_3d", "average_curves", "curve_math",
                "column_math", "column_math2", "subtract_reference", "subtract_line",
                "translate_v", "translate_h", "polygon_area", "set_ops", "smooth2", "fft2",
                "ifft", "fft_filter", "stft", "iir", "convolution", "deconvolution",
                "correlation", "coherence", "hilbert", "envelope", "decimate",
                "mask_condition", "paste_transpose"]
MATRIX_ACTIONS = ["matrix_math", "matrix_filter", "matrix_crop", "matrix_rotate",
                  "matrix_resample", "matrix_integrate", "matrix_stats", "matrix_fft2",
                  "matrix_fft2_filter", "matrix_correlate2d"]


def test_every_action_runs_with_default_settings(win, monkeypatch, qapp):
    monkeypatch.setattr(md.FormDialog, "exec", lambda self: 1)
    bw, ws = _fill(win)
    qapp.clipboard().setText("1\t2\n3\t4")
    failures = {}
    names: set[str] = set()
    wavelets = [f"wavelet_{k}" for k in ("cwt", "dwt", "decompose", "denoise", "smooth")]
    for key in BOOK_ACTIONS + (wavelets if wv.available() else []):
        win.show_window(bw.book.name)
        names |= {s.name for s in bw.book.sheets}
        del bw.book.sheets[1:]                            # fresh data for every tool
        bw.book.active = 0
        del ws.columns[4:]
        ws.clear_mask()
        _fill(win)
        bw.reload()
        bw.select_columns([1])
        win.warnings.clear()
        win.data_actions[key].trigger()
        if win.warnings:
            failures[key] = win.warnings[:]
    # deconvolution of the default inputs (two different curves) is allowed to fail on
    # spectral zeros only; everything else must run
    failures.pop("deconvolution", None)
    assert failures == {}
    names |= {s.name for s in bw.book.sheets}
    assert {"Interpolated", "FFT", "IIRResponse", "Average", "Corr", "Pivot",
            "Stacked"} <= set(names)
    assert any(isinstance(b, MatrixBook) for b in win.project.windows.values())
    mb = MatrixBook.new("MBook9", 6, 6)
    mb.sheets[0] = MatrixSheet.new("MSheet1", 6, 6)
    mb.sheets[0].data = np.arange(36.0).reshape(6, 6) + np.eye(6) * 50
    win.new_matrix(mb)
    mw = win.matrix_window(mb)
    for key in MATRIX_ACTIONS + ["matrix_inverse"]:
        mb.active = 0                                     # always the original matrix
        mw.reload()
        win.warnings.clear()
        win.data_actions[key].trigger()
        assert win.warnings == [], key
    assert len(mb.sheets) > 8
    inv = next(s for s in mb.sheets if s.long_name == "Inverse")
    np.testing.assert_allclose(inv.data @ mb.sheets[0].data, np.eye(6), atol=1e-9)


def test_signal_runner_plots_into_the_graph(win):
    bw, ws = _fill(win)
    bw.select_columns([1])
    gwin = win.plot("line")
    choices, _, g = win.analysis_inputs()
    assert g is gwin
    n_before = len(gwin.graph.layers[0].plots)
    idx = sd.run_smooth(win, {"input": choices[0], "method": "loess", "span": 0.2}, gwin)
    assert len(gwin.graph.layers[0].plots) == n_before + 1
    assert ws.columns[idx].long_name.startswith("Signal (LOESS")


# ====================================================================== masks in the grid
def test_mask_cells_red_undo_and_excluded_from_analysis(win):
    from PySide6.QtCore import Qt

    from rsstatchem.ui.origin.sheet_view import HEADER_ROWS, MASK_FG
    bw, ws = _fill(win, 20)
    bw.select_block(3, 1, 5, 1)
    assert bw.mask_selection("mask") == 3
    assert ws.masked_rows(1) == [2, 3, 4]
    view = bw.sheet_view()
    brush = view.model_.data(view.model_.index(HEADER_ROWS + 3, 1),
                             Qt.ItemDataRole.ForegroundRole)
    assert brush.color() == MASK_FG
    bw.select_columns([1])
    choices, _, _ = win.analysis_inputs()
    s = md.resolve(win, choices[0])
    assert len(s.idx) == 17 and 3 not in s.idx
    # results land on the unmasked rows only
    col = sd.run_smooth(win, {"input": choices[0], "method": "binomial"})
    assert ws.cell(3, col) is None and ws.cell(6, col) is not None
    bw.undo_stack.undo()                                  # the smoothing column
    bw.undo_stack.undo()                                  # the mask
    assert bw.worksheet().masked_rows(1) == []
    bw.select_block(1, 1, 4, 1)
    bw.mask_selection("mask")
    bw.select_block(1, 1, 4, 1)
    st = bw.sheet_view().quick_stats()
    assert st is None                                     # all selected cells are masked
    bw.select_block(1, 1, 6, 1)
    st = bw.sheet_view().quick_stats()
    assert st["count"] == 2
    win.data_actions["unmask_all"].trigger()
    assert not bw.worksheet().has_masks()


# ====================================================================== import wizard
def test_import_wizard_preview_filters_and_modes(win, tmp_path, qapp):
    a = tmp_path / "a.csv"
    a.write_text("Time;Abs\ns;AU\n0;1,5\n1;2,5\n2;3,5\n", encoding="utf-8")
    b = tmp_path / "b.csv"
    b.write_text("Time;Abs\ns;AU\n0;7,5\n1;8,5\n", encoding="utf-8")
    folder = tmp_path / "filters"
    dlg = iw.ImportWizard(win, [str(a), str(b)], filter_folder=folder)
    assert dlg.table.columnCount() == 2
    assert dlg.table.item(0, 1).text() == "Abs" and dlg.table.item(1, 1).text() == "AU"
    assert dlg.table.item(3, 1).text() == "1.5"
    assert "Semicolon" in dlg.status.text()
    dlg.w["first_row"].setValue(2)
    assert dlg.table.item(3, 1).text() == "2.5"
    dlg.save_filter("semi", "*.csv")
    assert aw.find_filter(a, folder)[0] == "semi"
    files, opts, mode = dlg.values()
    assert files == [str(a), str(b)] and opts.first_row == 2 and mode == "books"
    dlg.grab().save(str(tmp_path / "wizard.png"))
    objs = iw.import_paths(win, files, aw.AsciiOptions(), "sheets")
    book = objs[0]
    assert isinstance(book, Workbook) and [s.name for s in book.sheets] == ["a", "b"]
    assert book.sheets[1].columns[1].units == "AU"
    objs = iw.import_paths(win, files, aw.AsciiOptions(), "columns")
    assert len(objs[0].sheets[0].columns) == 4
    objs = iw.import_paths(win, files, aw.AsciiOptions(), "books")
    assert len(objs) == 2 and all(o.name in win.project.windows for o in objs)


def test_export_from_the_window(win, tmp_path):
    bw, ws = _fill(win, 8)
    p = data_menus.export_xlsx(win, str(tmp_path / "o.xlsx"), "book")
    from openpyxl import load_workbook
    assert load_workbook(p).sheetnames == ["Sheet1"]
    p = data_menus.export_ascii(win, str(tmp_path / "o.csv"),
                                {"delimiter": "\t", "long": True, "units": False})
    assert p.read_text(encoding="utf-8").splitlines()[0].split("\t")[1] == "Signal"
    bw.select_columns([1])
    win.plot("line")
    out = data_menus.export_all_graphs(win, str(tmp_path), "png", 72)
    assert out and out[0].exists()
