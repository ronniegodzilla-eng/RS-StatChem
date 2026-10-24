"""Output Viewer: SPSS-style copy to Word, portrait fitting, transpose and zoom."""
import pytest

from rsstatchem.core import report
from rsstatchem.core.output import Image, OutputDocument, Section, Table, Text

NARROW = Table("Group Statistics", ["", "N", "Mean"],
               [["Kontrol", "41", "66.237"], ["Eksperimen", "39", "72.213"]])
WIDE = Table("Independent Samples Test",
             ["", ""] + [("Levene's Test for Equality of Variances", s) for s in ("F", "Sig.")]
             + [("t-test for Equality of Means", s) for s in
                ("t", "df", "One-Sided p", "Two-Sided p", "Mean Difference",
                 "Std. Error Difference", "95% CI Lower", "95% CI Upper")],
             [["Nilai post-test", "Equal variances assumed", ".332", ".566", "-2.446", "78",
               ".008", ".017", "-5.9762", "2.4432", "-10.8403", "-1.1122"],
              [None, "Equal variances not assumed", None, None, "-2.444", "77.416", ".008",
               ".017", "-5.9762", "2.4454", "-10.8453", "-1.1072"]],
             row_header_cols=2, footnotes=["a. Footnote."])


def test_portrait_fitting_shrinks_only_wide_tables():
    assert report.fit_for_portrait(NARROW) == (9.0, False)
    size, stretch = report.fit_for_portrait(WIDE)
    assert size < 9.0 and stretch
    widths = report.column_widths_in(WIDE, size, stretch)
    assert len(widths) == len(WIDE.columns)
    assert sum(widths) == pytest.approx(report.PORTRAIT_WIDTH_IN, abs=0.01)
    narrow = report.column_widths_in(NARROW, 9.0, False)
    assert sum(narrow) < report.PORTRAIT_WIDTH_IN  # not spread over the page


def test_spanning_header_labels_get_room():
    table = Table("t", ["", ("Maximum", "Statistic")], [["a", "1"]])
    alone = Table("t", ["", "Statistic"], [["a", "1"]])
    assert report.column_widths_in(table, 9, False)[1] > \
        report.column_widths_in(alone, 9, False)[1]


def test_transpose_swaps_rows_and_columns():
    t = report.transpose(WIDE)
    assert t.title == WIDE.title and t.row_header_cols == 1 and t.footnotes == WIDE.footnotes
    assert t.columns[1:] == ["Nilai post-test Equal variances assumed",
                             "Nilai post-test Equal variances not assumed"]
    assert t.rows[0] == ["Levene's Test for Equality of Variances F", ".332", ""]
    assert t.rows[-1] == ["t-test for Equality of Means 95% CI Upper", "-1.1122", "-1.1072"]
    assert report.fit_for_portrait(t)[0] == 9.0  # the tall version fits at full size


def test_word_html_is_self_contained():
    html = report.clipboard_html([WIDE])
    assert html.startswith("<html>") and "<style" not in html  # inline styles only
    assert "border-top:1.5pt solid #000" in html and "width:100%" in html
    assert 'colspan="8"' in html and "a. Footnote." in html
    assert "margin:0" in html  # no Word paragraph spacing inside cells
    narrow = report.word_table_html(NARROW)
    assert "font-size:9pt" in narrow and "width:100%" not in narrow


def test_tsv_and_text_for_excel():
    tsv = report.table_tsv(NARROW).splitlines()
    assert tsv[0] == "Group Statistics"
    assert tsv[2].split("\t") == ["Kontrol", "41", "66.237"]
    assert "t-test for Equality of Means t" in report.table_tsv(WIDE)
    section = Section("T-Test", [NARROW, Text("", "note text")])
    assert report.item_text(section).startswith("T-Test\n\nGroup Statistics")


def test_docx_tables_have_fixed_widths(tmp_path):
    from docx import Document

    doc = OutputDocument()
    doc.add(Section("T-Test", [NARROW, WIDE]))
    path = tmp_path / "out.docx"
    report.export_docx(doc, str(path))
    tables = Document(str(path)).tables
    assert len(tables) == 2
    xml = tables[1]._tbl.xml
    assert 'w:type="fixed"' in xml and "w:cantSplit" in xml and "w:tblHeader" in xml
    assert tables[1].cell(2, 1).text == "Equal variances assumed"


# ------------------------------------------------------------------ viewer
@pytest.fixture
def viewer(qapp, tmp_path):
    from PySide6.QtCore import QSettings

    from rsstatchem.ui.output_viewer import OutputViewer

    settings = QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat)
    v = OutputViewer(settings)
    v.resize(900, 600)
    png = report_png()
    v.add_section(Section("T-Test", [Text("", "T-TEST GROUPS=x", style="log"), NARROW, WIDE,
                                     Image("Plot", png, 120, 80)]))
    return v


def report_png() -> bytes:
    from PySide6.QtCore import QBuffer, QByteArray, QIODevice
    from PySide6.QtGui import QImage

    image = QImage(12, 8, QImage.Format.Format_RGB32)
    image.fill(0xFF0000)
    data = QByteArray()
    buffer = QBuffer(data)
    buffer.open(QIODevice.OpenModeFlag.WriteOnly)
    image.save(buffer, "PNG")
    return bytes(data)


def test_copy_selected_table_puts_html_text_and_image_on_clipboard(viewer):
    from PySide6.QtGui import QGuiApplication

    viewer.select(f"i{WIDE.id}")
    assert "Independent Samples Test" in viewer.view.toPlainText()
    viewer.copy()
    mime = QGuiApplication.clipboard().mimeData()
    assert mime.hasHtml() and mime.hasText() and mime.hasImage()
    assert "Equal variances not assumed" in mime.html()
    assert mime.text().startswith("Independent Samples Test")
    viewer.copy("text")
    assert not QGuiApplication.clipboard().mimeData().hasHtml()


def test_click_selects_the_object_under_the_mouse(viewer):
    viewer.show()
    viewer.select(None)
    cursor = viewer.view.document().find("Eksperimen")
    rect = viewer.view.cursorRect(cursor)
    assert viewer._anchor_at(rect.center()) == f"i{NARROW.id}"


def test_transpose_replaces_table_and_keeps_tree_in_sync(viewer):
    viewer.select(f"i{WIDE.id}")
    viewer.transpose_selected()
    section = viewer.doc.sections[0]
    new = section.items[2]
    assert new is not WIDE and new.columns[1].startswith("Nilai post-test")
    assert viewer.selected == f"i{new.id}"
    assert viewer._tree_node(viewer.selected) is not None


def test_zoom_steps_scales_and_persists(viewer):
    viewer.show()
    base = viewer.view.document().idealWidth()
    viewer.step_zoom(1)
    assert viewer.zoom == 1.1
    viewer.set_zoom(2.0)
    assert viewer.view.document().idealWidth() > base * 1.5
    assert float(viewer.settings.value("output/zoom")) == 2.0
    viewer.set_zoom(99)
    assert viewer.zoom == 3.0
    viewer.act_zoom_reset.trigger()
    assert viewer.zoom == 1.0 and "100%" in viewer.zoom_label.text()


def test_wide_table_scrolls_instead_of_breaking_words(viewer):
    viewer.resize(500, 400)
    viewer.show()
    viewer.set_zoom(1.5)
    text = viewer.view.toPlainText()
    assert "Difference" in text
    assert viewer.view.horizontalScrollBar().maximum() > 0
