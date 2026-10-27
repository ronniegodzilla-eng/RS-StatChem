"""Export an Output Viewer document to Excel (.xlsx), like SPSS's *Export Output > Excel*.

Two layouts: every table on its own worksheet (charts on their own sheets too), or the whole
output on one worksheet in viewer order. Tables keep the SPSS look in Excel: bold title,
shaded bold header with merged spanning labels, horizontal rules at the top, under the header
and at the bottom, and footnotes underneath. Numeric cells are written as numbers with a number
format that shows the same decimals (".045" stays ".045"); text such as "<.001" stays text.
The APA TableLook writes "Table N" plus an italic title, an unshaded header and a *Note.*
"""
from __future__ import annotations

import io
import re

from rsstatchem.core.output import Image, OutputDocument, Table, Text
from rsstatchem.core.pivot import visible

_NUMBER = re.compile(r"^(-?)(\d*)(?:\.(\d+))?(%?)$")
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")


def parse_number(text) -> tuple[float, str] | None:
    """(value, Excel number format) for a formatted SPSS number, else None."""
    if text is None:
        return None
    s = str(text).strip()
    m = _NUMBER.match(s)
    if not m or not (m.group(2) or m.group(3)):
        return None
    sign, whole, dec, pct = m.groups()
    value = float(f"{sign}{whole or '0'}.{dec or '0'}")
    decimals = len(dec or "")
    if pct:
        return value / 100, "0" + ("." + "0" * decimals if decimals else "") + "%"
    if decimals == 0:
        return value, "0"
    lead = "0" if whole else "#"
    return value, f"{lead}.{'0' * decimals}"


def _sheet_name(title: str, taken: set[str]) -> str:
    base = _BAD_SHEET.sub(" ", title).strip()[:28] or "Sheet"
    name, k = base, 2
    while name.lower() in taken:
        suffix = f" ({k})"
        name = base[:31 - len(suffix)] + suffix
        k += 1
    taken.add(name.lower())
    return name


def _header_cells(table: Table) -> list[list[tuple[str, int, int, int]]]:
    from rsstatchem.core.report import _header_rows

    return _header_rows(table)


def write_table(ws, table: Table, row: int, look: str = "spss", number: int = 1,
                col: int = 1) -> int:
    """Write ``table`` at (row, col); returns the next free row."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    table = visible(table)
    apa = look == "apa"
    font_name = "Times New Roman" if apa else "Arial"
    thick, thin = Side(style="medium"), Side(style="thin")
    head_fill = None if apa else PatternFill("solid", fgColor="F2F2F2")
    if apa:
        ws.cell(row, col, f"Table {number}").font = Font(name=font_name, bold=True, size=11)
        row += 1
        ws.cell(row, col, table.title).font = Font(name=font_name, italic=True, size=11)
    else:
        ws.cell(row, col, table.title).font = Font(name=font_name, bold=True, size=11)
    row += 1
    n_cols = len(table.columns)
    rh = table.row_header_cols
    header = _header_cells(table)
    first_header = row
    for r, cells in enumerate(header):
        for text, c, colspan, rowspan in cells:
            cell = ws.cell(row + r, col + c, text)
            cell.font = Font(name=font_name, bold=not apa, size=9)
            cell.alignment = Alignment(horizontal="left" if c < rh else "center",
                                       vertical="bottom", wrap_text=True)
            if colspan > 1 or rowspan > 1:
                ws.merge_cells(start_row=row + r, start_column=col + c,
                               end_row=row + r + rowspan - 1, end_column=col + c + colspan - 1)
    last_header = row + len(header) - 1
    for r in range(first_header, last_header + 1):
        for c in range(n_cols):
            cell = ws.cell(r, col + c)
            top = thick if r == first_header else None
            bottom = thin if r == last_header else None
            if r < last_header and not apa:
                bottom = None
            cell.border = Border(top=top, bottom=bottom)
            if head_fill is not None:
                cell.fill = head_fill
    row = last_header + 1
    for r, values in enumerate(table.rows):
        last = r == len(table.rows) - 1
        for c in range(n_cols):
            text = values[c] if c < len(values) else ""
            text = "" if text is None else str(text)
            cell = ws.cell(row, col + c)
            parsed = parse_number(text) if c >= rh else None
            if parsed is not None:
                cell.value, cell.number_format = parsed
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = text if text != "" else None
                cell.alignment = Alignment(horizontal="left", vertical="top",
                                           wrap_text=c < rh and len(text) > 40)
            cell.font = Font(name=font_name, size=9)
            if last:
                cell.border = Border(bottom=thick)
            if c < rh and head_fill is not None and not apa:
                cell.fill = PatternFill("solid", fgColor="FAFAFA")
        row += 1
    for k, note in enumerate(table.footnotes):
        text = f"Note. {note}" if apa and k == 0 and not note.lower().startswith("note.") \
            else note
        ws.cell(row, col, text).font = Font(name=font_name, size=8, italic=apa and k == 0)
        row += 1
    _fit_widths(ws, table, col)
    return row + 1


def _fit_widths(ws, table: Table, col: int) -> None:
    from openpyxl.utils import get_column_letter

    for c in range(len(table.columns)):
        texts = [str(r[c]) for r in table.rows if c < len(r) and r[c] is not None]
        head = table.columns[c]
        head_words = " ".join(head) if isinstance(head, tuple) else str(head)
        longest = max([len(t) for t in texts] + [max((len(w) for w in head_words.split()),
                                                     default=4)] + [4])
        letter = get_column_letter(col + c)
        width = min(max(longest + 2, 6), 60)
        current = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(current, width)


def _write_image(ws, item: Image, row: int, col: int = 1) -> int:
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font

    ws.cell(row, col, item.title).font = Font(name="Arial", bold=True, size=10)
    row += 1
    try:
        img = XLImage(io.BytesIO(item.png))
    except Exception:  # noqa: BLE001 - Pillow missing or unreadable PNG: keep the title only
        return row + 1
    scale = min(1.0, item.width / max(img.width, 1))
    img.width, img.height = int(img.width * scale), int(img.height * scale)
    ws.add_image(img, ws.cell(row, col).coordinate)
    return row + int(img.height / 20) + 2


def export_xlsx(doc: OutputDocument, path: str, layout: str = "sheets",
                look: str | None = None, include_charts: bool = True,
                include_notes: bool = True) -> int:
    """Write the output to ``path``; ``layout`` is "sheets" (one worksheet per table) or
    "single" (everything on one sheet). Returns the number of tables written."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    look = look if look in ("spss", "apa") else "spss"
    wb = Workbook()
    taken: set[str] = set()
    n_tables = 0
    if layout == "single":
        ws = wb.active
        ws.title = _sheet_name("Output", taken)
        row = 1
        for section in doc.sections:
            ws.cell(row, 1, section.title).font = Font(name="Arial", bold=True, size=13)
            row += 2
            for item in section.items:
                if isinstance(item, Table):
                    n_tables += 1
                    row = write_table(ws, item, row, look, n_tables)
                elif isinstance(item, Image) and include_charts:
                    row = _write_image(ws, item, row)
                elif isinstance(item, Text) and include_notes and item.style != "log":
                    if item.title:
                        ws.cell(row, 1, item.title).font = Font(name="Arial", bold=True)
                        row += 1
                    for line in item.text.splitlines() or [""]:
                        ws.cell(row, 1, line)
                        row += 1
                    row += 1
    else:
        wb.remove(wb.active)
        notes_ws = None
        for section in doc.sections:
            for item in section.items:
                if isinstance(item, Table):
                    n_tables += 1
                    ws = wb.create_sheet(_sheet_name(f"{n_tables} {item.title}", taken))
                    ws.cell(1, 1, section.title).font = Font(name="Arial", italic=True,
                                                            size=9, color="666666")
                    write_table(ws, item, 2, look, n_tables)
                elif isinstance(item, Image) and include_charts:
                    ws = wb.create_sheet(_sheet_name(f"Chart {item.title}", taken))
                    _write_image(ws, item, 1)
                elif isinstance(item, Text) and include_notes and item.style != "log":
                    if notes_ws is None:
                        notes_ws = wb.create_sheet(_sheet_name("Notes", taken))
                        notes_ws.column_dimensions["A"].width = 100
                    r = notes_ws.max_row + (2 if notes_ws.max_row > 1 else 0)
                    notes_ws.cell(r, 1, f"{section.title}: {item.title}".rstrip(": ")).font = \
                        Font(bold=True)
                    for k, line in enumerate(item.text.splitlines() or [""], 1):
                        notes_ws.cell(r + k, 1, line)
        if not wb.worksheets:
            wb.create_sheet("Output")
    wb.save(path)
    return n_tables
