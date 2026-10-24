"""Render an :class:`OutputDocument` to HTML (viewer, HTML export, clipboard) and Word (.docx).

Three renderings of the same tables:

* the Output Viewer (Qt rich text): SPSS-like look, zoomable, the selected item framed;
* the clipboard for Word ("Copy", like SPSS): inline-styled HTML that Word turns into a real
  table, sized to fit a portrait page (smaller font, then wrap) when it would be too wide;
* Word export (.docx) with the same fitting rules.
"""
from __future__ import annotations

import base64
import html
import io
from typing import Callable

from rsstatchem import APP_NAME, __version__
from rsstatchem.core.output import Image, OutputDocument, Section, Table, Text

ACCENT = "#B71C1C"
FONT = "Arial, Helvetica, sans-serif"
PORTRAIT_WIDTH_IN = 6.3   # A4/Letter portrait text width with ~2.5 cm margins
_PAD_IN = 0.15            # left+right cell padding

# Arial/Helvetica advance widths (1/1000 em) for sizing columns the way Word will set them.
_DIGITS = dict.fromkeys("0123456789#$_?", 556)
_REGULAR = {**_DIGITS, " ": 278, ".": 278, ",": 278, ":": 278, ";": 278, "!": 278, "/": 278,
            "-": 333, "(": 333, ")": 333, "[": 278, "]": 278, "'": 191, '"': 355, "*": 389,
            "%": 889, "&": 667, "+": 584, "=": 584, "<": 584, ">": 584, "@": 1015,
            **dict(zip("ABCDEFGHIJKLMNOPQRSTUVWXYZ",
                       (667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722,
                        778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611))),
            **dict(zip("abcdefghijklmnopqrstuvwxyz",
                       (556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556,
                        556, 556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500)))}
_BOLD = {**_REGULAR, "'": 238, "J": 556, "K": 722, "L": 611, "A": 722, "B": 722,
         **dict(zip("abcdefghijklmnopqrstuvwxyz",
                    (556, 611, 556, 611, 556, 333, 611, 611, 278, 278, 556, 278, 889, 611,
                     611, 611, 611, 389, 556, 333, 611, 556, 778, 556, 556, 500)))}


def _text_em(text: str, bold: bool = False) -> float:
    widths = _BOLD if bold else _REGULAR
    return sum(widths.get(ch, 556) for ch in text) / 1000.0


def css(scale: float = 1.0) -> str:
    def pt(size: float) -> str:
        return f"{size * scale:.2f}pt"

    def px(size: float) -> str:  # Qt's rich text only honours px for padding/margins
        return f"{max(1, round(size * scale))}px"

    return f"""
body {{ font-family: {FONT}; font-size: {pt(9)}; color: #111; background: #ffffff; }}
p.section {{ font-size: {pt(12.5)}; color: #1b1b1b; font-weight: bold;
             margin-top: {px(18)}; margin-bottom: {px(4)}; }}
p.caption {{ font-size: {pt(9.5)}; font-weight: bold; margin-top: {px(12)};
             margin-bottom: {px(3)}; }}
table.rs {{ border-collapse: collapse; }}
table.rs th, table.rs td {{ padding-top: {px(2)}; padding-bottom: {px(2)};
                            padding-left: {px(7)}; padding-right: {px(7)}; font-size: {pt(9)}; }}
table.rs th {{ font-weight: bold; background: #f2f2f2; vertical-align: bottom; }}
table.rs td.num {{ white-space: nowrap; }}
p.foot {{ font-size: {pt(8)}; color: #333; margin: 0; }}
pre.log {{ font-family: Menlo, Consolas, monospace; font-size: {pt(8)}; color: #666;
           background: #f7f7f7; border-left: 3px solid {ACCENT}; padding: 3px 8px; }}
p.note {{ margin: {pt(3)} 0; }}
p.warning {{ margin: {pt(3)} 0; color: #8a5a00; background: #fff7e0; padding: 4px 8px; }}
"""


CSS = css()


def _esc(value) -> str:
    return html.escape(str(value))


def _looks_numeric(text: str) -> bool:
    t = text.strip().replace(",", "").replace("%", "").rstrip("*ᵃᵇᶜᵈᵉᶠᵍʰ")
    if t in ("", ".", "-"):
        return False
    try:
        float(t.lstrip("<>"))
        return True
    except ValueError:
        return False


# ------------------------------------------------------------------ table structure
def _header_rows(table: Table) -> list[list[tuple[str, int, int, int]]]:
    """Header as rows of (text, column index, colspan, rowspan)."""
    cols = list(table.columns)
    if not table.has_groups:
        return [[(str(c), i, 1, 1) for i, c in enumerate(cols)]]
    top, bottom = [], []
    i = 0
    while i < len(cols):
        c = cols[i]
        if isinstance(c, tuple):
            span = 1
            while i + span < len(cols) and isinstance(cols[i + span], tuple) \
                    and cols[i + span][0] == c[0]:
                span += 1
            top.append((c[0], i, span, 1))
            bottom.extend((cols[j][1], j, 1, 1) for j in range(i, i + span))
            i += span
        else:
            top.append((str(c), i, 1, 2))
            i += 1
    return [top, bottom]


def _column_ems(table: Table, wrap_row_headers: bool) -> list[float]:
    """Width (in em) each column needs: headers wrap at spaces, numbers never wrap, row-header
    text wraps at spaces only when ``wrap_row_headers`` (or when it is very long)."""
    def longest_word(text: str, bold: bool = False) -> float:
        return max([_text_em(w, bold) for w in text.split()] or [0.0])

    n = len(table.columns)
    widest = [0.0] * n
    for header in _header_rows(table):  # a spanning label shares its columns' width
        for text, col, colspan, _ in header:
            share = longest_word(text, bold=True) / colspan
            for c in range(col, col + colspan):
                widest[c] = max(widest[c], share)
    for row in table.rows:
        for c, cell in enumerate(row[:n]):
            text = "" if cell is None else str(cell)
            wraps = c < table.row_header_cols and (wrap_row_headers or len(text) > 24)
            widest[c] = max(widest[c], longest_word(text) if wraps else _text_em(text))
    return widest


def natural_width_in(table: Table, font_pt: float = 9.0, wrap_row_headers: bool = False,
                     pad_in: float = _PAD_IN) -> float:
    """Width of a table in inches at ``font_pt`` (headers wrap at spaces)."""
    return sum(em * font_pt / 72 + pad_in for em in _column_ems(table, wrap_row_headers))


def fit_for_portrait(table: Table) -> tuple[float, bool]:
    """(font size in pt, stretch to page width?) so the table fits a portrait page:
    shrink the font first, then let row labels wrap and spread over the page width."""
    for size in (9.0, 8.0, 7.5):
        if natural_width_in(table, size) <= PORTRAIT_WIDTH_IN:
            return size, False
    for size in (8.0, 7.5, 7.0):
        if natural_width_in(table, size, wrap_row_headers=True, pad_in=0.09) <= \
                PORTRAIT_WIDTH_IN:
            return size, True
    return 7.0, True


def transpose(table: Table) -> Table:
    """Swap rows and columns (SPSS: pivot), e.g. to turn a very wide table into a tall one
    that fits a portrait page. Two-level headers become "Group Sub" labels."""
    rh = table.row_header_cols
    labels = [" ".join(p for p in c if p) if isinstance(c, tuple) else str(c)
              for c in table.columns]
    # each original row becomes a column, titled by its (filled-down) row-header cells
    filled, last = [], [""] * rh
    for row in table.rows:
        heads = [str(row[i]) if row[i] not in (None, "") else "" for i in range(rh)]
        heads = [h or (last[i] if all(not heads[j] for j in range(i + 1)) else "")
                 for i, h in enumerate(heads)] if rh else []
        last = [h or last[i] for i, h in enumerate(heads)]
        filled.append(" ".join(h for h in heads if h).strip() or "(row)")
    corner = " / ".join(l for l in labels[:rh] if l)
    columns = [corner] + filled
    rows = [[labels[c]] + [("" if r[c] is None else str(r[c])) for r in table.rows]
            for c in range(rh, len(table.columns))]
    return Table(table.title, columns, rows, row_header_cols=1,
                 footnotes=list(table.footnotes))


# ------------------------------------------------------------------ viewer HTML
_TOP = "border-top: 1.5px solid #000;"
_UNDER = "border-bottom: 1px solid #000;"
_BOTTOM = "border-bottom: 1.5px solid #000;"


def table_html(table: Table, anchor: str = "") -> str:
    out = [f'<a name="{anchor}"></a>' if anchor else "",
           f'<p class="caption">{_esc(table.title)}</p>',
           '<table class="rs" cellspacing="0">', "<thead>"]
    rh = table.row_header_cols
    rows = _header_rows(table)
    for r, row in enumerate(rows):
        cells = []
        for text, col, colspan, rowspan in row:
            style = ""
            if r == 0:
                style += _TOP
            if r == len(rows) - 1 or rowspan == len(rows):
                style += _UNDER
            elif colspan > 1:
                style += "border-bottom: 1px solid #999;"
            align = "left" if col < rh else "center"
            span = (f' colspan="{colspan}"' if colspan > 1 else "") + \
                (f' rowspan="{rowspan}"' if rowspan > 1 else "")
            cells.append(f'<th{span} align="{align}" style="{style}">{_esc(text)}</th>')
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</thead><tbody>")
    for r, row in enumerate(table.rows):
        style = f' style="{_BOTTOM}"' if r == len(table.rows) - 1 else ""
        cells = []
        for i, cell in enumerate(row):
            text = "" if cell is None else str(cell)
            if i < rh:
                cells.append(f'<td class="rh" align="left"{style}>{_esc(text)}</td>')
            else:
                align = "right" if _looks_numeric(text) else "left"
                cells.append(f'<td class="num" align="{align}"{style}>{_esc(text)}</td>')
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</tbody></table>")
    for note in table.footnotes:
        out.append(f'<p class="foot">{_esc(note)}</p>')
    return "\n".join(out)


def _png_data_uri(img: Image) -> str:
    return "data:image/png;base64," + base64.b64encode(img.png).decode("ascii")


def item_html(item, image_src: Callable[[Image], str] = _png_data_uri, scale: float = 1.0) -> str:
    anchor = f"i{item.id}"
    if isinstance(item, Table):
        return table_html(item, anchor)
    if isinstance(item, Text):
        if item.style == "log":
            return f'<a name="{anchor}"></a><pre class="log">{_esc(item.text)}</pre>'
        title = f'<p class="caption">{_esc(item.title)}</p>' if item.title else ""
        return (f'<a name="{anchor}"></a>{title}<p class="{item.style}">'
                f'{_esc(item.text).replace(chr(10), "<br>")}</p>')
    if isinstance(item, Image):
        return (f'<a name="{anchor}"></a><p class="caption">{_esc(item.title)}</p>'
                f'<p><img src="{image_src(item)}" width="{int(item.width * scale)}"></p>')
    return ""


def _framed(content: str) -> str:
    """The selected item, framed like SPSS's selection box."""
    return ('<table cellspacing="0" cellpadding="4" border="1" '
            f'style="border: 2px solid {ACCENT}; border-color: {ACCENT};"><tr><td>'
            f"{content}</td></tr></table>")


def section_html(section: Section, image_src: Callable[[Image], str] = _png_data_uri,
                 scale: float = 1.0, selected: str | None = None) -> str:
    parts = [f'<a name="s{section.id}"></a><p class="section">{_esc(section.title)}</p>']
    for item in section.items:
        block = item_html(item, image_src, scale)
        parts.append(_framed(block) if selected == f"i{item.id}" else block)
    body = "\n".join(parts)
    return _framed(body) if selected == f"s{section.id}" else body


def document_html(doc: OutputDocument, image_src: Callable[[Image], str] = _png_data_uri,
                  standalone: bool = True, scale: float = 1.0,
                  selected: str | None = None) -> str:
    body = "\n<br>\n".join(section_html(s, image_src, scale, selected) for s in doc.sections)
    style = css(scale)
    if not standalone:
        return f"<html><head><style>{style}</style></head><body>{body}</body></html>"
    return ("<!doctype html><html><head><meta charset='utf-8'>"
            f"<title>{APP_NAME} Output</title><style>{style}</style></head>"
            f"<body>{body}</body></html>")


# ------------------------------------------------------------------ clipboard (Word / Excel)
def _cell_p(text: str) -> str:
    """Cell text as a spacing-free paragraph; Word otherwise adds its Normal style's space
    after every cell and the rows come out double-spaced."""
    return f'<p style="margin:0; line-height:normal;">{_esc(text)}</p>'


def word_table_html(table: Table) -> str:
    """A table as Word-friendly HTML: inline styles only, fitted to a portrait page."""
    size, stretch = fit_for_portrait(table)
    pad = "1pt 3pt" if stretch else "1pt 5pt"
    base = f"font-family:{FONT}; font-size:{size:g}pt; padding:{pad};"
    width = "width:100%;" if stretch else ""
    out = [f'<p style="font-family:{FONT}; font-size:{size + 1:g}pt; font-weight:bold; '
           f'margin:0 0 3pt 0;">{_esc(table.title)}</p>',
           f'<table cellspacing="0" cellpadding="0" style="border-collapse:collapse; {width}">']
    rh = table.row_header_cols
    rows = _header_rows(table)
    for r, row in enumerate(rows):
        cells = []
        for text, col, colspan, rowspan in row:
            border = ""
            if r == 0:
                border += "border-top:1.5pt solid #000;"
            if r == len(rows) - 1 or rowspan == len(rows):
                border += "border-bottom:1pt solid #000;"
            elif colspan > 1:
                border += "border-bottom:0.5pt solid #000;"
            align = "left" if col < rh else "center"
            span = (f' colspan="{colspan}"' if colspan > 1 else "") + \
                (f' rowspan="{rowspan}"' if rowspan > 1 else "")
            cells.append(f'<td{span} style="{base} {border} font-weight:bold; '
                         f'text-align:{align}; vertical-align:bottom;">{_cell_p(text)}</td>')
        out.append("<tr>" + "".join(cells) + "</tr>")
    for r, row in enumerate(table.rows):
        bottom = "border-bottom:1.5pt solid #000;" if r == len(table.rows) - 1 else ""
        cells = []
        for i, cell in enumerate(row):
            text = "" if cell is None else str(cell)
            if i < rh:
                cells.append(f'<td style="{base} {bottom} text-align:left; '
                             f'vertical-align:top;">{_cell_p(text)}</td>')
            else:
                align = "right" if _looks_numeric(text) else "left"
                cells.append(f'<td style="{base} {bottom} text-align:{align}; '
                             f'vertical-align:top; white-space:nowrap;">{_cell_p(text)}</td>')
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</table>")
    for note in table.footnotes:
        out.append(f'<p style="font-family:{FONT}; font-size:{max(size - 1.5, 7):g}pt; '
                   f'margin:0;">{_esc(note)}</p>')
    return "\n".join(out)


def word_item_html(item) -> str:
    if isinstance(item, Table):
        return word_table_html(item)
    if isinstance(item, Text):
        if item.style == "log":
            return (f'<p style="font-family:Consolas, monospace; font-size:8pt; color:#555;">'
                    f'{_esc(item.text).replace(chr(10), "<br>")}</p>')
        title = (f'<p style="font-family:{FONT}; font-size:10pt; font-weight:bold; '
                 f'margin:6pt 0 2pt 0;">{_esc(item.title)}</p>') if item.title else ""
        return (f'{title}<p style="font-family:{FONT}; font-size:10pt; margin:0 0 4pt 0;">'
                f'{_esc(item.text).replace(chr(10), "<br>")}</p>')
    if isinstance(item, Image):
        width_px = min(item.width, int(PORTRAIT_WIDTH_IN * 96))
        return (f'<p style="font-family:{FONT}; font-size:10pt; font-weight:bold;">'
                f'{_esc(item.title)}</p><p><img src="{_png_data_uri(item)}" '
                f'width="{width_px}"></p>')
    return ""


def clipboard_html(items: list) -> str:
    """Complete HTML document for the clipboard (Qt wraps it as CF_HTML on Windows)."""
    parts = []
    for item in items:
        if isinstance(item, Section):
            parts.append(f'<p style="font-family:{FONT}; font-size:12pt; font-weight:bold; '
                         f'margin:8pt 0 4pt 0;">{_esc(item.title)}</p>')
            parts.extend(word_item_html(i) for i in item.items
                         if getattr(i, "style", "") != "log")
        else:
            parts.append(word_item_html(item))
    body = "\n<p style='margin:0'>&nbsp;</p>\n".join(parts)
    return f"<html><head><meta charset='utf-8'></head><body>{body}</body></html>"


def table_tsv(table: Table) -> str:
    """Tab-separated text (pastes into Excel cell by cell)."""
    header = [" ".join(c) if isinstance(c, tuple) else str(c) for c in table.columns]
    lines = [table.title, "\t".join(header)]
    lines += ["\t".join("" if v is None else str(v) for v in row) for row in table.rows]
    lines += table.footnotes
    return "\n".join(lines)


def item_text(item) -> str:
    if isinstance(item, Table):
        return table_tsv(item)
    if isinstance(item, Text):
        return (f"{item.title}\n" if item.title else "") + item.text
    if isinstance(item, Image):
        return item.title
    if isinstance(item, Section):
        return "\n\n".join([item.title] + [item_text(i) for i in item.items])
    return ""


# ---------------------------------------------------------------- Word export
def _set_cell_border(cell, **edges) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, size in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if size else "nil")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), "000000")
        borders.append(el)


def column_widths_in(table: Table, size: float, stretch: bool) -> list[float]:
    """Column widths (inches) for Word, as wide as the content needs. When that is wider
    than a portrait page, row labels wrap at spaces and get whatever room the numbers
    leave; only a table that still overflows is squeezed."""
    def inches(ems: list[float], pad: float) -> list[float]:
        return [em * size / 72 * 1.05 + pad for em in ems]

    full = inches(_column_ems(table, wrap_row_headers=False), _PAD_IN)
    if not stretch and sum(full) <= PORTRAIT_WIDTH_IN:
        return full
    rh = table.row_header_cols
    widths = inches(_column_ems(table, wrap_row_headers=True), 0.1)
    spare = PORTRAIT_WIDTH_IN - sum(widths)
    for c in range(rh):
        grow = min(max(spare, 0.0), max(full[c] - widths[c], 0.0))
        widths[c] += grow
        spare -= grow
    total = sum(widths)
    if total > PORTRAIT_WIDTH_IN:
        widths = [w * PORTRAIT_WIDTH_IN / total for w in widths]
    return widths


def _table_layout(t, widths: list[float]) -> None:
    """Fixed column widths (Word's autofit would spread every table over the page) and
    tight cell margins."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    t.autofit = False
    for column, w in zip(t.columns, widths):
        column.width = Inches(w)
        for cell in column.cells:
            cell.width = Inches(w)
    tbl_pr = t._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(round(sum(widths) * 1440)))
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "70")
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tbl_pr.append(margins)


def _row_flags(row, header: bool) -> None:
    """Never split a row over two pages; header rows repeat on every page."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    for tag in ("w:cantSplit", "w:tblHeader") if header else ("w:cantSplit",):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "true")
        tr_pr.append(el)


def _cell_style(document, size: float):
    """Paragraph style for table cells at ``size`` pt: no spacing, single lines. Empty cells
    take their height from it too, so they no longer stretch rows to the 10 pt body size."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    name = f"RS Table {size:g}pt"
    if name in [s.name for s in document.styles]:
        return document.styles[name]
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.paragraph_format.space_before = Pt(1)
    style.paragraph_format.space_after = Pt(1)
    style.paragraph_format.line_spacing = 1.0
    return style


def _docx_table(document, table: Table) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    size, stretch = fit_for_portrait(table)
    cell_style = _cell_style(document, size)
    cap = document.add_paragraph()
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_after = Pt(3)
    run = cap.add_run(table.title)
    run.bold = True
    run.font.size = Pt(size + 1)

    header_rows = 2 if table.has_groups else 1
    n_cols = len(table.columns)
    t = document.add_table(rows=header_rows + len(table.rows), cols=n_cols)
    _table_layout(t, column_widths_in(table, size, stretch))  # before merging cells

    for r, header in enumerate(_header_rows(table)):
        for text, col, colspan, rowspan in header:
            cell = t.cell(r, col)
            if colspan > 1 or rowspan > 1:
                cell = cell.merge(t.cell(r + rowspan - 1, col + colspan - 1))
            cell.text = text
    for r, row in enumerate(table.rows):
        for c, value in enumerate(row[:n_cols]):
            t.cell(header_rows + r, c).text = "" if value is None else str(value)

    last = header_rows + len(table.rows) - 1
    for r, row in enumerate(t.rows):
        _row_flags(row, header=r < header_rows)
        for c, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.style = cell_style
                # the caption, header and first data row never get separated
                p.paragraph_format.keep_with_next = r <= header_rows and r < last
                if c >= table.row_header_cols:
                    numeric = r >= header_rows and _looks_numeric(cell.text)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r < header_rows else (
                        WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT)
                for run in p.runs:
                    run.bold = r < header_rows
            edges = {"left": 0, "right": 0}
            if r == 0:
                edges["top"] = 12
            if r == header_rows - 1:
                edges["bottom"] = 8
            if r == last:
                edges["bottom"] = 12
            _set_cell_border(cell, **edges)

    for note in table.footnotes:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(note)
        run.font.size = Pt(max(size - 1.5, 7))
    document.add_paragraph()


def export_docx(doc: OutputDocument, path: str) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    for section in doc.sections:
        heading = document.add_heading(section.title, level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0x1B, 0x1B, 0x1B)
            run.font.name = "Arial"
            run.font.size = Pt(13)
        for item in section.items:
            if isinstance(item, Table):
                _docx_table(document, item)
            elif isinstance(item, Text):
                if item.style == "log":
                    continue
                if item.title:
                    document.add_paragraph().add_run(item.title).bold = True
                document.add_paragraph(item.text)
            elif isinstance(item, Image):
                p = document.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.add_run(item.title).bold = True
                document.add_picture(io.BytesIO(item.png),
                                     width=Inches(min(item.width / 96, PORTRAIT_WIDTH_IN)))
    core = document.core_properties
    core.comments = f"Generated by {APP_NAME} {__version__}"
    document.save(path)
