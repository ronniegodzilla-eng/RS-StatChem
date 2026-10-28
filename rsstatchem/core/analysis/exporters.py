"""Export of worksheets and report sheets (Origin's File > Export), without Qt.

Worksheets are duck-typed: any object with ``name``, ``columns`` (each with ``name``,
``long_name``, ``units``, ``comments``, ``values``) and ``data_rows()`` — the Graph &
Analysis :class:`~rsstatchem.ui.origin.model.Worksheet` fits.

* ASCII / CSV with selectable header rows (short name, Long Name, Units, Comments),
  delimiter and decimal separator; masked cells can be written blank.
* Excel ``.xlsx`` with one worksheet per sheet (openpyxl), header rows in bold.
* Report sheets to HTML (from the caller's HTML) or Word ``.docx`` (python-docx): tables,
  notes and plot images.
"""
from __future__ import annotations

import base64
import csv
import io
import math
from pathlib import Path

HEADER_ROWS = {"name": "Short Name", "long_name": "Long Name", "units": "Units",
               "comments": "Comments"}


def _cell_text(v, decimal: str = ".") -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        s = str(int(v)) if v.is_integer() and abs(v) < 1e15 else f"{v:.15g}"
        return s.replace(".", ",") if decimal == "," else s
    return str(v)


def header_lines(cols, headers) -> list[list[str]]:
    """Header rows: the short name and Long Name rows always (Long Name falls back to the
    short name), Units / Comments only when some column has one."""
    out = []
    for h in headers:
        line = [getattr(c, h, "") or "" for c in cols]
        if h in ("name", "long_name") or any(line):
            if h == "long_name":
                line = [c.long_name or c.name for c in cols]
            out.append(line)
    return out


def sheet_table(ws, headers=("long_name", "units"), masked: str = "keep",
                columns=None) -> list[list]:
    """Rows to export: the chosen header rows (only those with any text, except the short
    name) followed by the data rows. ``masked="blank"`` empties masked cells."""
    cols = [ws.columns[i] for i in (range(len(ws.columns)) if columns is None else columns)]
    rows: list[list] = header_lines(cols, headers)
    n = ws.data_rows()
    for r in range(n):
        line = []
        for c in cols:
            v = c.values[r] if r < len(c.values) else None
            if masked == "blank" and r in getattr(c, "mask", ()):
                v = None
            line.append(v)
        rows.append(line)
    return rows


def write_csv(ws, path: str | Path, delimiter: str = ",", decimal: str = ".",
              headers=("long_name", "units"), masked: str = "keep", encoding: str = "utf-8") \
        -> Path:
    p = Path(path)
    rows = sheet_table(ws, headers, masked)
    with open(p, "w", newline="", encoding=encoding) as fh:
        w = csv.writer(fh, delimiter=delimiter)
        for r in rows:
            w.writerow([_cell_text(v, decimal) for v in r])
    return p


def write_xlsx(sheets, path: str | Path, headers=("long_name", "units", "comments"),
               masked: str = "keep") -> Path:
    """``sheets``: list of ``(sheet name, worksheet)``; every worksheet becomes an Excel
    sheet (names made unique and cut to Excel's 31 characters)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for name, ws in sheets:
        base = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "Sheet"
        title, k = base, 2
        while title.lower() in used:
            suffix = f"_{k}"
            title, k = base[:31 - len(suffix)] + suffix, k + 1
        used.add(title.lower())
        xs = wb.create_sheet(title)
        n_head = len(header_lines(ws.columns, headers))
        for i, row in enumerate(sheet_table(ws, headers, masked)):
            xs.append([None if v is None or (isinstance(v, float) and math.isnan(v)) else v
                       for v in row])
            if i < n_head:
                for cell in xs[i + 1]:
                    cell.font = Font(bold=True)
        xs.freeze_panes = xs.cell(row=n_head + 1, column=1)
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    p = Path(path)
    wb.save(p)
    return p


def write_html(html: str, path: str | Path) -> Path:
    p = Path(path)
    p.write_text(html, encoding="utf-8")
    return p


def write_report_docx(report, path: str | Path, png_for=None) -> Path:
    """A report sheet (``title`` and ``items`` of kind table / text / image / graph) as a
    Word document. ``png_for(item)`` renders graph items (PNG bytes) — the UI passes the
    report view's renderer."""
    from docx import Document
    from docx.shared import Cm
    doc = Document()
    doc.add_heading(getattr(report, "title", "") or getattr(report, "name", "Report"), 1)
    group = None
    for item in report.items:
        g = item.get("group", "")
        if g and g != group:
            doc.add_heading(g, 2)
        group = g
        kind = item.get("kind")
        if item.get("title"):
            doc.add_heading(item["title"], 3 if g else 2)
        if kind == "table":
            cols = [c if isinstance(c, str) else " ".join(map(str, c))
                    for c in item.get("columns", [])]
            rows = item.get("rows", [])
            width = max([len(cols)] + [len(r) for r in rows] + [1])
            t = doc.add_table(rows=1 + len(rows), cols=width)
            t.style = "Table Grid"
            for j, c in enumerate(cols):
                t.cell(0, j).text = str(c)
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    t.cell(i + 1, j).text = str(v)
            for note in item.get("notes", []):
                doc.add_paragraph(str(note)).italic = True
        elif kind == "text":
            doc.add_paragraph(item.get("text", ""))
        elif kind in ("image", "graph"):
            png = None
            if item.get("png"):
                png = base64.b64decode(item["png"])
            elif png_for is not None:
                png = png_for(item)
            if png:
                doc.add_picture(io.BytesIO(png), width=Cm(15))
    p = Path(path)
    doc.save(p)
    return p
