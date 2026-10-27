"""Reading and writing datasets: RS-StatChem project (.rsd), SPSS (.sav, .por), Stata (.dta),
SAS (.sas7bdat, .xpt), CSV/TXT and Excel. Statistical formats go through pyreadstat.

A ``.rsd`` file is a zip archive with ``meta.json`` (variable dictionary) and ``data.csv``
(values as text, decoded per variable type on load), so it stays readable by other tools.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

from rsstatchem import __version__
from rsstatchem.core.dataset import (DATE, MEASURES, NOMINAL, NUMERIC, SCALE, STRING, Dataset,
                                     Variable, format_code)

PROJECT_EXT = ".rsd"
PROJECT_FORMAT = "rs-statchem-dataset"

OPEN_FILTER = ("All supported (*.rsd *.sav *.csv *.txt *.tsv *.dat *.xlsx *.xlsm *.rsg *.rsc "
               "*.rso *.dta *.sas7bdat *.xpt *.por);;"
               "RS-StatChem data (*.rsd);;SPSS data (*.sav);;"
               "Text / CSV (*.csv *.txt *.tsv *.dat);;Excel (*.xlsx *.xlsm);;"
               "Graph (*.rsg);;Chemical drawing (*.rsc);;RS-StatChem output (*.rso);;"
               "Stata (*.dta);;SAS (*.sas7bdat *.xpt);;SPSS portable (*.por)")
STATA_FILTER = "Stata (*.dta)"
SAS_FILTER = "SAS (*.sas7bdat *.xpt)"
POR_FILTER = "SPSS portable (*.por)"


# ------------------------------------------------------------------ project
def save_project(ds: Dataset, path: str | Path) -> None:
    meta = {"format": PROJECT_FORMAT, "version": 1, "app_version": __version__,
            "name": ds.name, "variables": [v.to_dict() for v in ds.variables],
            "meta": ds.meta}
    df = ds.df.copy()
    for v in ds.variables:
        if v.type == DATE:
            df[v.name] = df[v.name].dt.strftime("%Y-%m-%dT%H:%M:%S")
    buf = io.StringIO()
    df.to_csv(buf, index=False, lineterminator="\n")
    tmp = Path(str(path) + ".tmp")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("meta.json", json.dumps(meta, ensure_ascii=False, indent=1))
        zf.writestr("data.csv", buf.getvalue())
    tmp.replace(path)


def load_project(path: str | Path) -> Dataset:
    with zipfile.ZipFile(path) as zf:
        meta = json.loads(zf.read("meta.json").decode("utf-8"))
        if meta.get("format") != PROJECT_FORMAT:
            raise ValueError("This file is not an RS-StatChem dataset.")
        raw = pd.read_csv(io.BytesIO(zf.read("data.csv")), dtype=str, keep_default_na=False,
                          na_filter=False)
    variables = [Variable.from_dict(d) for d in meta["variables"]]
    data = {}
    for v in variables:
        col = raw[v.name] if v.name in raw else pd.Series([""] * len(raw))
        if v.type == NUMERIC:
            data[v.name] = pd.to_numeric(col.replace("", np.nan), errors="coerce")
            v.value_labels = {float(k): lbl for k, lbl in v.value_labels.items()}
            v.missing = [float(m) for m in v.missing]
        elif v.type == DATE:
            data[v.name] = pd.to_datetime(col.replace("", None), errors="coerce")
        else:
            data[v.name] = col.astype(object)
    ds = Dataset(pd.DataFrame(data, index=pd.RangeIndex(len(raw))), variables,
                 name=meta.get("name", "DataSet1"))
    ds.meta = meta.get("meta", {}) or {}
    ds.path = str(path)
    return ds


# ------------------------------------------------------------------ text / csv
ENCODINGS = ("utf-8", "utf-8-sig", "cp1252", "latin-1")


def detect_encoding(path: str | Path) -> str:
    sample = Path(path).read_bytes()[:200_000]
    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    for enc in ("utf-8", "cp1252"):
        try:
            sample.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def sniff_text(path: str | Path, encoding: str) -> tuple[str, str]:
    """Guess (delimiter, decimal separator). Handles the common ';' + ',' European/ID style."""
    with open(path, encoding=encoding, errors="replace", newline="") as fh:
        sample = fh.read(64_000)
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = max(",;\t|", key=sample.count)
    decimal = "."
    if delimiter != ",":
        comma_numbers = len(re.findall(r"(?<![\w.,])-?\d+,\d+(?![\w.,])", sample))
        dot_numbers = len(re.findall(r"(?<![\w.,])-?\d+\.\d+(?![\w.,])", sample))
        if comma_numbers > dot_numbers:
            decimal = ","
    return delimiter, decimal


def read_text(path: str | Path, delimiter: str = ",", decimal: str = ".", header: bool = True,
              encoding: str = "utf-8", nrows: int | None = None,
              skip_rows: int = 0) -> pd.DataFrame:
    df = pd.read_csv(path, sep=delimiter, decimal=decimal, header=0 if header else None,
                     encoding=encoding, nrows=nrows, skiprows=skip_rows, skipinitialspace=True,
                     thousands=None)
    if not header:
        df.columns = [f"VAR{i + 1:05d}" for i in range(df.shape[1])]
    return _clean_frame(df)


def _clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    df = df.dropna(axis=1, how="all") if df.shape[1] > 1 else df
    unnamed = [c for c in df.columns if str(c).startswith("Unnamed:") and df[c].isna().all()]
    return df.drop(columns=unnamed)


# ------------------------------------------------------------------ excel
def excel_sheets(path: str | Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_excel(path: str | Path, sheet: str | int = 0, header: bool = True) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet, header=0 if header else None, engine="openpyxl")
    if not header:
        df.columns = [f"VAR{i + 1:05d}" for i in range(df.shape[1])]
    return _clean_frame(df.dropna(axis=0, how="all"))


def write_excel(ds: Dataset, path: str | Path, use_labels: bool = False) -> None:
    df = ds.df.copy()
    if use_labels:
        for v in ds.variables:
            if v.value_labels:
                df[v.name] = df[v.name].map(lambda x, v=v: v.value_labels.get(x, x))
    dictionary = pd.DataFrame([{
        "Name": v.name, "Type": v.type, "Label": v.label, "Measure": v.measure,
        "Decimals": v.decimals,
        "Values": "; ".join(f"{format_code(k)}={lbl}" for k, lbl in v.value_labels.items()),
        "Missing": ", ".join(format_code(m) for m in v.missing)} for v in ds.variables])
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Data", index=False)
        dictionary.to_excel(writer, sheet_name="Variables", index=False)


def write_text(ds: Dataset, path: str | Path, delimiter: str = ",", decimal: str = ".") -> None:
    ds.df.to_csv(path, sep=delimiter, decimal=decimal, index=False, date_format="%Y-%m-%d")


# ------------------------------------------------------------------ SPSS
_FMT_RE = re.compile(r"^([A-Z]+)(\d+)?(?:\.(\d+))?$")


_STATA_FMT = re.compile(r"^%-?(\d+)(?:\.(\d+))?([a-z]+)")
_SAS_FMT = re.compile(r"^([A-Z$]*?)(\d+)?\.(\d+)?$")


def read_sav(path: str | Path) -> Dataset:
    import pyreadstat

    df, meta = pyreadstat.read_sav(str(path), user_missing=True, dates_as_pandas_datetime=True)
    return _from_readstat(df, meta, path)


def read_por(path: str | Path) -> Dataset:
    import pyreadstat

    df, meta = pyreadstat.read_por(str(path), dates_as_pandas_datetime=True)
    return _from_readstat(df, meta, path)


def read_stata(path: str | Path) -> Dataset:
    import pyreadstat

    df, meta = pyreadstat.read_dta(str(path), dates_as_pandas_datetime=True)
    return _from_readstat(df, meta, path, "stata")


def read_sas(path: str | Path) -> Dataset:
    """SAS data (.sas7bdat) or SAS transport (.xpt)."""
    import pyreadstat

    if Path(path).suffix.lower() == ".xpt":
        df, meta = pyreadstat.read_xport(str(path), dates_as_pandas_datetime=True)
    else:
        df, meta = pyreadstat.read_sas7bdat(str(path), dates_as_pandas_datetime=True)
    return _from_readstat(df, meta, path, "sas")


def _format_parts(fmt: str | None, source: str) -> tuple[str, str | None, str | None]:
    """(SPSS-like kind, width, decimals) from a SPSS/Stata/SAS display format."""
    fmt = fmt or ""
    if source == "stata":
        m = _STATA_FMT.match(fmt)
        if not m:
            return "F", "8", "2"
        width, dec, kind = m.groups()
        if kind.startswith("t"):
            return "DATE", width, None
        return ("A" if kind.endswith("s") else "F"), width, dec
    if source == "sas":
        m = _SAS_FMT.match(fmt.upper())
        if not m:
            return "F", "8", None
        kind, width, dec = m.groups()
        if kind.startswith("$"):
            return "A", width, None
        if any(k in kind for k in ("DATE", "DDMM", "MMDD", "YYMM", "TIME")):
            return "DATE", width, None
        return "F", width or "8", dec
    m = _FMT_RE.match(fmt)
    return (m.group(1), m.group(2), m.group(3)) if m else ("F", "8", "2")


def _from_readstat(df: pd.DataFrame, meta, path: str | Path, source: str = "spss") -> Dataset:
    labels = dict(zip(meta.column_names, meta.column_labels or [None] * len(meta.column_names)))
    variables = []
    for name in meta.column_names:
        fmt = (meta.original_variable_types or {}).get(name, "F8.2")
        kind, width, dec = _format_parts(fmt, source)
        series = df[name]
        if kind == "A" or series.dtype == object and not pd.api.types.is_numeric_dtype(series):
            vtype = STRING
        elif "DATE" in kind or "TIME" in kind or pd.api.types.is_datetime64_any_dtype(series):
            vtype = DATE
        else:
            vtype = NUMERIC
        measure = (meta.variable_measure or {}).get(name, "unknown")
        if measure not in MEASURES:
            measure = NOMINAL if vtype == STRING else SCALE
        value_labels = (meta.variable_value_labels or {}).get(name, {}) or {}
        missing = []
        for rng in (meta.missing_ranges or {}).get(name, []) or []:
            lo, hi = rng.get("lo"), rng.get("hi")
            if lo == hi and lo is not None:
                missing.append(lo)
        if vtype == NUMERIC and dec is None and source != "spss":
            values = pd.to_numeric(series, errors="coerce").dropna()
            dec = "0" if len(values) == 0 or bool(np.all(np.mod(values, 1) == 0)) else "2"
        if vtype == NUMERIC:
            value_labels = {float(k): str(v) for k, v in value_labels.items()}
            missing = [float(x) for x in missing]
        else:
            value_labels = {str(k): str(v) for k, v in value_labels.items()}
            missing = [str(x) for x in missing]
        align = (getattr(meta, "variable_alignment", None) or {}).get(name)
        variables.append(Variable(
            name=name, type=vtype, width=int(width or 8),
            decimals=int(dec or 0) if vtype == NUMERIC else 0,
            label=labels.get(name) or "", value_labels=value_labels, missing=missing,
            columns=int((getattr(meta, "variable_display_width", None) or {}).get(name) or 8),
            align=align if align in ("left", "right", "center")
            else ("left" if vtype == STRING else "right"),
            measure=measure))
    ds = Dataset(df, variables, name=Path(path).stem)
    ds.path = str(path)
    return ds


def write_sav(ds: Dataset, path: str | Path) -> None:
    import pyreadstat

    df = ds.df.copy()
    formats, value_labels, missing = {}, {}, {}
    for v in ds.variables:
        if v.type == STRING:
            df[v.name] = df[v.name].fillna("").astype(str)
            width = max(v.width, int(df[v.name].str.len().max() or 1))
            formats[v.name] = f"A{width}"
        elif v.type == NUMERIC:
            formats[v.name] = f"F{max(v.width, v.decimals + 2)}.{v.decimals}"
        else:
            formats[v.name] = "DATE11"
        if v.value_labels:
            value_labels[v.name] = dict(v.value_labels)
        if v.missing:
            missing[v.name] = list(v.missing)[:3]
    pyreadstat.write_sav(
        df, str(path),
        column_labels={v.name: (v.label or None) for v in ds.variables},
        variable_value_labels=value_labels or None,
        missing_ranges=missing or None,
        variable_display_width={v.name: v.columns for v in ds.variables},
        variable_measure={v.name: v.measure for v in ds.variables},
        variable_format=formats)


def _export_frame(ds: Dataset, max_name: int, pattern: str) -> tuple[pd.DataFrame, dict]:
    """Copy of the data with names valid for Stata/SAS (letters, digits, underscore)."""
    rename, used = {}, set()
    for v in ds.variables:
        base = re.sub(pattern, "_", v.name)[:max_name] or "v"
        if base[0].isdigit():
            base = ("v" + base)[:max_name]
        name, k = base, 1
        while name.lower() in used:
            suffix = f"_{k}"
            name = base[:max_name - len(suffix)] + suffix
            k += 1
        used.add(name.lower())
        rename[v.name] = name
    df = ds.df.copy()
    for v in ds.variables:
        if v.type == STRING:
            df[v.name] = df[v.name].fillna("").astype(str)
    return df.rename(columns=rename), rename


def write_stata(ds: Dataset, path: str | Path) -> dict:
    """Stata 15 .dta with variable and (integer) value labels; returns the name mapping."""
    import pyreadstat

    df, rename = _export_frame(ds, 32, r"[^A-Za-z0-9_]")
    labels, formats = {}, {}
    for v in ds.variables:
        name = rename[v.name]
        if v.type == NUMERIC and v.value_labels and \
                all(float(k).is_integer() for k in v.value_labels):
            labels[name] = {int(k): str(t) for k, t in v.value_labels.items()}
        if v.type == NUMERIC:
            formats[name] = f"%{max(v.width, v.decimals + 2)}.{v.decimals}f"
    pyreadstat.write_dta(df, str(path), column_labels=[v.label or "" for v in ds.variables],
                         variable_value_labels=labels or None, variable_format=formats or None)
    return rename


def write_xport(ds: Dataset, path: str | Path) -> dict:
    """SAS transport file (XPORT version 8); returns the name mapping."""
    import pyreadstat

    df, rename = _export_frame(ds, 32, r"[^A-Za-z0-9_]")
    pyreadstat.write_xport(df, str(path), column_labels=[(v.label or "")[:40]
                                                         for v in ds.variables],
                           table_name=re.sub(r"[^A-Za-z0-9_]", "_", ds.name)[:32] or "DATA",
                           file_format_version=8)
    return rename


# ------------------------------------------------------------------ dispatch
def kind_of(path: str | Path) -> str:
    ext = Path(path).suffix.lower()
    if ext == PROJECT_EXT:
        return "project"
    if ext == ".sav":
        return "sav"
    if ext == ".por":
        return "por"
    if ext == ".dta":
        return "stata"
    if ext in (".sas7bdat", ".xpt"):
        return "sas"
    if ext == ".rso":
        return "output"
    if ext in (".xlsx", ".xlsm"):
        return "excel"
    if ext in (".csv", ".txt", ".tsv", ".dat"):
        return "text"
    raise ValueError(f"Unsupported file type: {ext or '(none)'}")
